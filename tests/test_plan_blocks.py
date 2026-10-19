"""Таблица плана блоками («2026 РЕКОМЕНДАЦИЯ» / «2026 ФАКТ») → длинный план."""
import io

import openpyxl
import pandas as pd

from atlas.engine.core.loader import load_file
from atlas.engine.core import plan_blocks, tabular
from atlas.imports import Imports

MONTHS = ['Апрель', 'Май', 'Июнь', 'Июль', 'Август', 'Сентябрь', 'Октябрь']


def book(tmp_path, name='plan.xlsx'):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.cell(2, 21).value = 'ГСП 1'                       # шапка диаграммы выше таблицы
    def block(r, title, rows):
        ws.cell(r, 20).value = title
        for j, m in enumerate(MONTHS): ws.cell(r + 1, 21 + j).value = m
        for i, (g, vals) in enumerate(rows):
            ws.cell(r + 2 + i, 20).value = g
            for j, v in enumerate(vals): ws.cell(r + 2 + i, 21 + j).value = v
        ws.cell(r + 2 + len(rows), 21).value = 120     # итоговая строка без названия
    block(8, '2026 РЕКОМЕНДАЦИЯ', [('ГСП 1', [30, 160, 195, 205, 200, 190, 95]), ('ГСП-2', [30, 166, 200, 200, 200, 195, None]),
                                  ('ГСП 8', [0, 115, 235, 245, 245, 220, 90])])
    block(20, '2026 ФАКТ', [('ГСП 1', [0, 256, 143, 181, 174, 124, None])])
    path = tmp_path / name; wb.save(path); return path


def test_blocks_to_long_plan(tmp_path):
    r = load_file(str(book(tmp_path)))
    plan = r.frames['plan']
    assert len(plan) == 20                                    # 7 + 6 (пустая пропущена) + 7; факт не в плане
    assert set(plan.group) == {'ГСП 1', 'ГСП 2', 'ГСП 8'}      # «ГСП-2» нормирована
    assert (plan.date.dt.day == 1).all() and plan.date.min() == pd.Timestamp('2026-04-01')
    assert plan[(plan.group == 'ГСП 8') & (plan.date == '2026-04-01')].plan_volume.iloc[0] == 0   # ноль остаётся
    assert set(plan.kind) == {'injection'}
    assert any('ФАКТ' in i['Причина'] for i in r.issues)     # о факте остаётся заметка


def test_winter_months_are_withdrawal():
    rows = [[None, '2026 ПЛАН', None, None], [None, 'Ноябрь', 'Декабрь', 'Январь'], ['ГСП 3', 5, 6, 7]]
    out, _ = plan_blocks.parse(rows)
    assert [r[2] for r in out] == ['отбор'] * 3
    assert [r[1] for r in out] == [pd.Timestamp('2026-11-01'), pd.Timestamp('2026-12-01'), pd.Timestamp('2027-01-01')]


def test_excel_download(tmp_path):
    p = book(tmp_path)
    imports = Imports.__new__(Imports)
    imports.file = lambda token: type('U', (), {'content': p.read_bytes(), 'name': p.name})()
    content, name = imports.plan_xlsx('x')
    frame = pd.read_excel(io.BytesIO(content))
    assert list(frame.columns) == list(plan_blocks.HEAD) and len(frame) == 20 and name.endswith('.xlsx')
    assert frame.iloc[0].tolist()[0] == 'ГСП 1' and frame.iloc[0]['Режим'] == 'закачка'


def test_plain_sheets_untouched(tmp_path):
    p = tmp_path / 'a.xlsx'; wb = openpyxl.Workbook(); wb.active.append(['Группа', 'Месяц', 'Режим', 'План, млн м³'])
    wb.active.append(['ГСП 9', '01.05.2025', 'закачка', 150]); wb.save(p)
    assert load_file(str(p)).frames['plan'].plan_volume.iloc[0] == 150
