"""ГИС: Excel «Σ hi × Кг» — формулы вместо значений, листы и совпадение с расчётом Атласа.
Данные — литеральные; реальных данных компании в репозитории нет. Проверка пересчётом — только если есть LibreOffice."""
import io
import shutil
import subprocess

import openpyxl
import pandas as pd
import pytest

from atlas import gis_xlsx
from atlas.engine.modules import gis_hikg


def _prepared():
    raw = pd.DataFrame({
        'well': ['115'] * 4 + ['113'],
        'date': pd.to_datetime(['2024-10-01'] * 3 + ['2025-03-01'] + ['2024-10-01']),
        'season': ['осень'] * 3 + ['весна', 'осень'],
        'investigation': [1, 1, 1, 3, 4],
        'pack': ['верхняя', 'верхняя', 'нижняя', 'верхняя', 'верхняя'],
        'from_m': [1, 2, 3, 4, 5], 'to_m': [2, 3, 4, 5, 6],
        'hi_m': [2.0, 1.5, 4.0, 2.5, 1.0], 'kg': [0.30, 0.25, 0.20, 0.40, None],
    })
    return gis_hikg.prepare(raw)


def test_workbook_has_three_sheets_and_formulas_not_numbers():
    book = openpyxl.load_workbook(io.BytesIO(gis_xlsx.build(_prepared())))
    assert book.sheetnames == ['Интервалы', 'Исследования', 'Итог']
    assert str(book['Интервалы']['J2'].value).startswith('=IF(')          # hi × Кг — формула, не число
    assert str(book['Исследования']['G2'].value).startswith('=SUMIFS(')
    assert str(book['Итог']['E2'].value).startswith('=SUMIFS(')


def test_total_by_pack_headers_are_the_packs_present():
    book = openpyxl.load_workbook(io.BytesIO(gis_xlsx.build(_prepared())))
    header = [c.value for c in book['Итог'][1]]
    assert header[4:6] == ['верхняя', 'нижняя'] and header[-1] == 'Всего, Σ hi × Кг'


def test_empty_selection_is_a_clear_error():
    with pytest.raises(ValueError, match='нет интервалов'):
        gis_xlsx.build(_prepared().iloc[0:0])


@pytest.mark.skipif(shutil.which('soffice') is None, reason='нужен LibreOffice для пересчёта формул')
def test_recalculated_values_match_engine(tmp_path):
    src = tmp_path / 'hikg.xlsx'
    src.write_bytes(gis_xlsx.build(_prepared()))
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', str(tmp_path / 'r'),
                    str(src)], capture_output=True, timeout=100, check=False)
    values = openpyxl.load_workbook(tmp_path / 'r' / 'hikg.xlsx', data_only=True)['Итог']
    rows = {(r[0], r[3]): r for r in values.iter_rows(min_row=2, values_only=True)}
    engine = gis_hikg.sums(_prepared())
    for well, inv in (('115', '1'), ('115', '3'), ('113', '4')):
        expected = engine[(engine['well'] == well) & (engine['investigation'].astype(str) == inv)]['hikg'].sum()
        assert rows[(well, inv)][-1] == pytest.approx(expected)
    assert rows[('115', '1')][4] == pytest.approx(2.0 * 0.30 + 1.5 * 0.25)   # верхняя пачка исследования 1
