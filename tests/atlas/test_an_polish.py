"""Авторский надзор: устойчивость к плохим данным, названия без номеров, графики отбора, XLSX графика и таблиц."""
import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

from atlas import render
from atlas.contract import Axis, Chart, Column, Data, Series, Table
from atlas.domain import DatasetKind
from atlas.engine.core.loader import load_file
from atlas.modules import an_annulus, an_report

PRODUCTION = pd.DataFrame([
    ('1', '2025-11-01', 100_000), ('2', '2025-11-01', 50_000), ('1', '2025-11-02', 120_000), ('2', '2025-11-02', 0),
    ('1', '2025-12-01', 90_000), ('2', '2025-12-01', 40_000),
], columns=['well', 'date', 'q'])
PRODUCTION['date'] = pd.to_datetime(PRODUCTION.date)
PRODUCTION['kind'], PRODUCTION['period'] = 'withdrawal', '2025-2026'
PRESSURE = pd.DataFrame({'date': pd.to_datetime(['2025-11-01', '2025-12-01']), 'pressure': [90.0, 85.0]})
FACTOR = pd.DataFrame({'date': pd.to_datetime(['2025-11-01', '2025-11-02', '2025-12-01']), 'gas_in_place': [5e6, 4.9e6, 4.7e6],
                       'water_day': [10.0, 12.0, np.nan], 'wf': [0.5, 0.6, np.nan]})


def run(extra, params):
    data = Data({DatasetKind.PRODUCTION: PRODUCTION, DatasetKind.OBJECT_PRESSURE: PRESSURE, **extra})
    data.mapping = {'1': {'group': 'ГСП-1'}, '2': {'group': 'ГСП 2'}}
    return an_report.AuthorReportModule().run(data, {'kind': 'withdrawal', 'periods': ['2025-2026'], 'annulus': [], **params})


def test_withdrawal_season_builds_charts_from_input_data():
    result = run({DatasetKind.WATER_FACTOR: FACTOR}, {})
    ids = {c.id for c in result.charts}
    assert {'an-canvas-withdrawal-2025-2026', 'an-fund-withdrawal-2025-2026'} <= ids
    wells = next(s for c in result.charts if c.id.startswith('an-canvas') for s in c.series if s.name == 'Скважин в работе')
    assert list(wells.y) == [2, 1, 2]
    full = run({DatasetKind.WATER_FACTOR: FACTOR}, {'show': [v for _, v in an_report.an_charts.SHOW]})
    names = {s.name for c in full.charts if c.id.startswith('an-canvas') for s in c.series}
    assert {'Вынос воды', 'Водный фактор, суточный', 'Объём газа в пласте', 'Накопленный отбор'} <= names
    assert not any(c.title.lower().startswith(('табл', 'рис')) or ' табл. ' in c.title for c in result.charts)
    assert not any(t.title.lower().startswith(('табл', 'рис')) for t in result.tables)
    assert not any('раздел в разработке' in n.text for n in result.notes)


def test_bad_well_number_in_annulus_is_a_note_not_a_crash():
    rows = pd.DataFrame({'well': ['78', '`', '81'], 'date': pd.Timestamp('2025-09-01'), 'season': 'Межсезонье 2025',
                         'q': [50.4, 3.0, 96.5], 'p': [9.0, 1.0, 35.0]})
    result = run({DatasetKind.ANNULUS: rows}, {'annulus': ['Межсезонье 2025']})
    table = next(t for t in result.tables if t.id == 'an-annulus-113')
    assert table.title == 'Результаты замеров МКД'
    assert table.frame['№ скв.'].tolist() == [78, 81]
    assert any('«`»' in n.text and n.level == 'warning' for n in result.notes)
    assert {'an-annulus-p_max', 'an-annulus-q_max'} <= {c.id for c in result.charts}
    assert an_annulus.valid_rows(rows)[1] == ['`']


def test_object_pressure_two_column_file_is_recognised(tmp_path):
    path = tmp_path / 'p.csv'
    path.write_text('Дата;Давление, кгс/см²\n02.08.2006;94,7\n04.08.2006;94,9\n', encoding='utf-8')
    r = load_file(path, 'auto', 'withdrawal', 'м³/сут', 'тыс. м³/сут')
    assert list(r.frames) == ['object_pressure']
    assert r.frames['object_pressure'].pressure.tolist() == [94.7, 94.9]


def sample_chart():
    chart = Chart('t', 'Расход и давление', Axis('Дата', scale='time'), Axis('Расход', 'тыс. м³/сут'), y2=Axis('Рпл', 'кгс/см²'))
    days = pd.date_range('2025-11-01', periods=4).to_numpy()
    chart.series.append(Series('Расход', days, np.array([1.0, 2.0, np.nan, 4.0]), 'line', labels=list('abcd')))
    chart.series.append(Series('Рпл', days[:2], np.array([90.0, 89.0]), 'points', axis='y2'))
    return chart


def test_chart_xlsx_has_all_points_and_an_excel_chart():
    content, name, mime = render.chart_file(sample_chart(), 'xlsx', 300)
    assert name == 'Расход и давление.xlsx' and 'spreadsheetml' in mime
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ['Данные', 'График']
    rows = list(wb['Данные'].values)
    assert rows[0][0] == 'Расход · Дата' and 'Расход · Расход, тыс. м³/сут' in rows[0] and 'Рпл · Рпл, кгс/см²' in rows[0]
    assert len(rows) == 5 and rows[3][1] is None                      # пропуск остаётся пустой ячейкой
    assert len(wb['График']._charts) == 1


def test_category_chart_xlsx_is_stacked():
    chart = Chart('c', 'Группы', Axis('Группа', scale='category', categories=['ГСП 1', 'ГСП 2']), Axis('Q', 'тыс. м³/сут'))
    chart.series.append(Series('Ноябрь', ['ГСП 1', 'ГСП 2'], [1.0, 2.0], 'bar', stack='g'))
    chart.series.append(Series('Декабрь', ['ГСП 1', 'ГСП 2'], [3.0, np.nan], 'bar', stack='g'))
    wb = openpyxl.load_workbook(io.BytesIO(render.chart_file(chart, 'xlsx', 300)[0]))
    assert [list(r) for r in wb['Данные'].values] == [['Группа', 'Ноябрь, тыс. м³/сут', 'Декабрь, тыс. м³/сут'],
                                                      ['ГСП 1', 1.0, 3.0], ['ГСП 2', 2.0, None]]


def test_tables_xlsx_is_formatted_like_the_screen():
    frame = pd.DataFrame({'a': ['ГСП 1', 'ГСП 2'], 'q': [1.234, 5.0], 'd': pd.to_datetime(['2025-11-01', '2025-11-02'])})
    table = Table('t', 'Таблица средних', frame, [Column('a', 'Группа'), Column('q', 'Расход', 'тыс. м³/сут', 1, 'number'),
                                                  Column('d', 'Дата', kind='date')], note='Примечание')
    wb = openpyxl.load_workbook(io.BytesIO(render.tables_file([table], 'Средние')[0]))
    ws = wb['Таблица средних']
    assert ws['B1'].value == 'Расход, тыс. м³/сут' and ws['B1'].font.bold and ws['B2'].number_format == '0.0'
    assert ws['C2'].number_format == 'DD.MM.YYYY' and ws.freeze_panes == 'A2'
    assert ws['A5'].value == 'Примечание'


def test_three_axes_export_to_all_formats():
    from atlas import render
    result = run({}, {})
    chart = next(c for c in result.charts if c.id.startswith('an-canvas'))
    assert chart.y3 is not None
    for fmt in ('svg', 'png', 'pdf', 'xlsx'):
        data, name, _ = render.chart_file(chart, fmt, 150)
        assert data and name.endswith(fmt)
    fig = render.to_plotly(chart)
    assert {t.yaxis for t in fig.data} == {'y', 'y2', 'y3'} and fig.layout.yaxis3.overlaying == 'y'
