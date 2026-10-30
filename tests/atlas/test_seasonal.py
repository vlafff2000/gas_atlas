"""«Сезонные показатели»: объём сезона, источник (узел учёта / скважины), Δ, группы."""
import pandas as pd
import pytest

from atlas.contract import Data
from atlas.domain import DatasetKind
from atlas.modules import _seasonal as S
from atlas.modules.seasonal import SeasonalModule

PROD, GAS = DatasetKind.PRODUCTION, DatasetKind.GAS_DB


def frame():
    rows = []
    for year, kind, period, start, rate_a, rate_b in (
            (2022, 'injection', '2022', '2022-05-01', 100000, 50000), (2022, 'withdrawal', '2022-2023', '2022-11-01', 200000, 100000),
            (2023, 'injection', '2023', '2023-05-01', 150000, 50000)):
        for d in pd.date_range(start, periods=10):
            rows.append({'well': '1', 'date': d, 'q': float(rate_a), 'kind': kind, 'period': period})
            rows.append({'well': '2', 'date': d, 'q': float(rate_b), 'kind': kind, 'period': period})
    rows.append({'well': '1', 'date': pd.Timestamp('2022-08-01'), 'q': 9e6, 'kind': 'injection', 'period': 'Нейтральный период Весна 2022'})
    return pd.DataFrame(rows)


def test_volume_days_and_delta():
    t = S.season_table(frame())
    inj = t[t.kind.eq('injection')].set_index('period')
    assert inj.loc['2022', 'volume'] == pytest.approx(1.5)               # (100+50) тыс. м³ × 10 сут
    assert inj.loc['2022', 'days'] == 10 and inj.loc['2022', 'mean'] == pytest.approx(150.0)
    assert inj.loc['2023', 'delta'] == pytest.approx(0.5) and inj.loc['2023', 'delta_pct'] == pytest.approx(100 / 3)
    assert 'Нейтральный период Весна 2022' not in set(t.period)
    c = S.cycles_table(t)
    # цикл = отбор + следующая за ним закачка: закачка 2022 — до отбора (цикл 2021), отбор 2022-2023 и закачка 2023 — цикл 2022
    assert list(c.cycle) == [2021, 2022] and pd.isna(c.w_volume.iloc[0]) and c.w_volume.iloc[1] == pytest.approx(3.0)
    assert c.i_volume.iloc[0] == pytest.approx(1.5) and c.i_volume.iloc[1] == pytest.approx(2.0)


def test_reference_source_and_discrepancy():
    ref = pd.DataFrame({'date': pd.date_range('2022-05-01', periods=10), 'kind': 'injection', 'gas': 165.0})   # тыс. м³/сут
    t = S.season_table(frame(), ref).set_index(['kind', 'period'])
    row = t.loc[('injection', '2022')]
    assert row.source == 'Узел учёта' and row.volume == pytest.approx(1.65) and row.diff_pct == pytest.approx(10.0)
    assert t.loc[('injection', '2023'), 'source'] == 'Сумма по скважинам'   # эталона за сезон нет
    forced = S.season_table(frame(), ref, S.SOURCE_WELLS).set_index(['kind', 'period'])
    assert forced.loc[('injection', '2022'), 'volume'] == pytest.approx(1.5)


def test_excluded_points_and_groups():
    df = frame()
    df['_excluded'] = (df.well == '2')
    assert S.season_table(df).set_index(['kind', 'period']).loc[('injection', '2022'), 'volume'] == pytest.approx(1.0)
    t = S.season_table(frame())
    g = S.group_table(frame(), t, {'1': 'ГСП 1', '2': 'ГСП 2'})
    one = g[(g.kind == 'injection') & (g.period == '2022')].set_index('group')
    assert one.loc['ГСП 1', 'volume'] == pytest.approx(1.0) and one.loc['ГСП 1', 'share'] == pytest.approx(100 * 2 / 3)


def test_module_run():
    data = Data({PROD: frame()})
    data.mapping = {'1': {'group': 'ГСП 1'}, '2': {'group': 'ГСП 2'}}
    r = SeasonalModule().run(data, {'source': 'auto', 'periods': [], 'groups': []})
    assert {t.id for t in r.tables} >= {'seasonal-cycles', 'seasonal-sources', 'seasonal-by-group'}
    assert r.charts[0].series[0].name == 'Закачка'
    empty = SeasonalModule().run(Data({PROD: frame().iloc[0:0]}), {'source': 'auto', 'periods': [], 'groups': []})
    assert not empty.tables and empty.notes


def test_excel_formulas_for_delta_and_discrepancy():
    import io
    import openpyxl
    from atlas.render import tables_file
    data = Data({PROD: frame()})
    data.mapping = {}
    tables = SeasonalModule().run(data, {}).tables
    book = openpyxl.load_workbook(io.BytesIO(tables_file(tables, 'x')[0]))
    ws = book['Источник объёма и расхождение'[:27]]
    head = [c.value for c in ws[1]]
    assert ws.cell(2, head.index('Δ к прошлому сезону, млн м³') + 1).value is None           # первый сезон
    delta = ws.cell(3, head.index('Δ к прошлому сезону, млн м³') + 1).value
    assert delta.startswith('=IF(A3=A2,') and 'D3-D2' in delta


def test_months_partial_delta_and_excel():
    import io
    import openpyxl
    from atlas.render import tables_file
    t = S.season_table(frame())
    m = S.month_table(frame(), t).set_index(['kind', 'period'])
    row = m.loc[('injection', '2023')]
    assert row.volume == pytest.approx(2.0) and row.delta == pytest.approx(0.5) and bool(row.partial)   # сезон из 10 суток мая — неполный месяц
    data = Data({PROD: frame()})
    data.mapping = {'1': {'group': 'ГСП 1'}, '2': {'group': 'ГСП 2'}}
    r = SeasonalModule().run(data, {'month_periods': ['2022', '2023']})
    assert {'seasonal-months', 'seasonal-months-groups'} <= {x.id for x in r.tables}
    assert {'seasonal-months-injection', 'seasonal-months-groups-injection'} <= {c.id for c in r.charts}
    bars = next(c for c in r.charts if c.id == 'seasonal-months-injection')
    assert [s.name for s in bars.series] == ['2022', '2023']
    ws = openpyxl.load_workbook(io.BytesIO(tables_file(r.tables, 'x')[0]))['Месяцы сезонов']
    head = [c.value for c in ws[1]]
    f = ws.cell(3, head.index('Δ к тому же месяцу прошлого сезона, млн м³') + 1).value
    assert f.startswith('=IF(COUNTIFS($A$2:$A$4,A3,') and 'SUMIFS($H$2:$H$4' in f
