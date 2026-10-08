import json

import pytest
from starlette.testclient import TestClient

from atlas.engine.core.storage import Store
from atlas.api import create_app
from atlas.projects import Projects


@pytest.fixture()
def env(tmp_path):
    projects = Projects(tmp_path)
    client = TestClient(create_app(projects))
    pid = client.post('/api/projects/demo').json()['id']
    return client, pid, Store(tmp_path)


def run(client, pid, **params):
    r = client.post('/api/modules/gdi/run', json={'project': pid, 'params': params})
    assert r.status_code == 200, r.text
    return r.json()


def table(body, tid):
    return next(t for t in body['tables'] if t['id'] == tid)


def test_health_and_modules(env):
    client, _, _ = env
    assert client.get('/api/health').json()['modules'] >= 1
    gdi = next(s for s in client.get('/api/modules').json() if s['id'] == 'gdi')
    assert gdi['needs'] == ['gdi']
    params = {p['name']: p for p in gdi['params']}
    assert params['wells']['source'] == {'dataset': 'gdi', 'column': 'well'}
    assert params['threshold']['setting'] == 'r2_threshold' and params['last_n']['default'] == 3


def test_project_listing_and_options(env):
    client, pid, _ = env
    listed = client.get('/api/projects').json()
    assert listed[0]['id'] == pid and listed[0]['demo'] is True and listed[0]['settings']['r2_threshold'] == .95
    wells = client.get(f'/api/projects/{pid}/options', params={'dataset': 'gdi', 'column': 'well'}).json()
    assert wells == ['31', '45', '70', '73', '89', '132', '540', '541']   # естественный порядок
    body = run(client, pid, wells=['31', '45'])
    assert table(body, 'studies')['count'] == 6 and len(body['charts']) == 2 and 'elapsed_ms' in body


def test_click_exclusion_undo_and_58_journal(env):
    client, pid, store = env
    body = run(client, pid, wells=['31'], last_n=1)
    point = next(s for s in body['charts'][0]['series'] if s['ids'])
    target = point['ids'][0]
    fit_before = table(body, 'studies')['rows'][[c['key'] for c in table(body, 'studies')['columns']].index('fit_points')][0]

    r = client.post(f'/api/projects/{pid}/exclusions',
                    json={'dataset': 'gdi', 'add': [target], 'reason': 'Исключено кликом на графике'})
    assert r.status_code == 200 and r.json()['added'] == 1
    m = store.manifest(pid)
    entry = m['settings']['excluded_points'][target]
    assert entry['module'] == 'gdi' and entry['metric'] == 'q' and entry['well'] == '31' and entry['batch']
    event = store.history(pid).iloc[0]
    details = json.loads(event['Подробности'])
    assert event['Действие'] == 'Ручной фильтр точек' and details['added_ids'] == [target] and details['gdi_after']

    body = run(client, pid, wells=['31'], last_n=1)
    cols = [c['key'] for c in table(body, 'studies')['columns']]
    assert table(body, 'studies')['rows'][cols.index('fit_points')][0] == fit_before - 1
    assert any(s['hollow'] for s in body['charts'][0]['series'])
    points = table(body, 'points')
    assert points['action']['checked'][points['action']['ids'].index(target)] is True

    assert client.post(f'/api/projects/{pid}/exclusions/undo', json={}).json()['removed'] == 1
    assert store.manifest(pid)['settings']['excluded_points'] == {}
    assert not any(s['hollow'] for s in run(client, pid, wells=['31'], last_n=1)['charts'][0]['series'])


def test_manual_filter_restores_unchecked(env):
    client, pid, store = env
    ids = table(run(client, pid, wells=['45'], last_n=1), 'points')['action']['ids'][:2]
    client.post(f'/api/projects/{pid}/exclusions', json={'dataset': 'gdi', 'add': ids, 'reason': 'Ручная проверка'})
    assert len(store.manifest(pid)['settings']['excluded_points']) == 2
    client.post(f'/api/projects/{pid}/exclusions', json={'dataset': 'gdi', 'remove': ids[:1]})
    assert list(store.manifest(pid)['settings']['excluded_points']) == ids[1:]
    r = client.post(f'/api/projects/{pid}/exclusions', json={'dataset': 'gdi', 'add': ['gdi:нет-такой']})
    assert r.status_code == 400


def test_threshold_is_shared_project_setting(env):
    client, pid, store = env
    r = client.patch(f'/api/projects/{pid}/settings', json={'values': {'r2_threshold': .9}})
    assert r.status_code == 200 and r.json()['settings']['r2_threshold'] == .9
    assert store.manifest(pid)['settings']['r2_threshold'] == .9
    assert client.patch(f'/api/projects/{pid}/settings', json={'values': {'r2_threshold': 2}}).status_code == 400
    assert client.patch(f'/api/projects/{pid}/settings', json={'values': {'excluded_points': {}}}).status_code == 400


def test_saved_view_and_history_in_58_format(env):
    client, pid, store = env
    assert client.get(f'/api/projects/{pid}/state/gdi').json() == {'panel': None, 'history': []}
    r = client.post(f'/api/projects/{pid}/state/gdi', json={'params': {'wells': ['31'], 'last_n': 2}})
    assert r.status_code == 200, r.text
    panel = store.manifest(pid)['settings']['panels']['gdi']
    assert panel['wells'] == ['31'] and panel['n'] == 2          # так читает 5.8
    saved = client.get(f'/api/projects/{pid}/state/gdi').json()
    assert saved['panel']['last_n'] == 2 and saved['history'][0]['params']['wells'] == ['31']
    event = store.history(pid)
    details = json.loads(event[event['Действие'] == 'Расчет ГДИ'].iloc[0]['Подробности'])
    assert details['wells'] == ['31'] and details['n'] == 2      # «Открыть расчет ГДИ» в 5.8


@pytest.mark.parametrize('fmt,magic', [('svg', b'<?xml'), ('png', b'\x89PNG'), ('pdf', b'%PDF')])
def test_chart_export(env, fmt, magic):
    client, pid, _ = env
    params = {'wells': ['31']}
    run(client, pid, **params)
    r = client.post('/api/modules/gdi/export', json={'project': pid, 'params': params, 'target': 'chart',
                                                     'id': 'gdi-31', 'format': fmt, 'dpi': 300})
    assert r.status_code == 200, r.text
    assert r.content.startswith(magic) and 'attachment' in r.headers['content-disposition']
    bad = client.post('/api/modules/gdi/export', json={'project': pid, 'params': params, 'target': 'chart',
                                                       'id': 'gdi-31', 'format': fmt, 'dpi': 7})
    assert bad.status_code == 400


def test_tables_export_xlsx(env):
    from io import BytesIO
    from openpyxl import load_workbook
    client, pid, _ = env
    r = client.post('/api/modules/gdi/export', json={'project': pid, 'params': {}, 'target': 'tables'})
    assert r.status_code == 200
    book = load_workbook(BytesIO(r.content), read_only=True)
    assert len(book.sheetnames) >= 3
    header = [c.value for c in next(book[book.sheetnames[0]].iter_rows(max_row=1))]
    assert header[:2] == ['Скважина', 'Дата'] and 'Qmax наблюд., тыс. м³/сут' in header
    one = client.post('/api/modules/gdi/export', json={'project': pid, 'params': {}, 'target': 'tables', 'id': 'studies'})
    assert len(load_workbook(BytesIO(one.content), read_only=True).sheetnames) == 1


def test_errors_are_readable(env):
    client, pid, _ = env
    r = client.post('/api/modules/gdi/run', json={'project': pid, 'params': {'threshold': 5}})
    assert r.status_code == 400 and 'Порог R²' in r.json()['error']
    assert client.post('/api/modules/nope/run', json={}).status_code == 404
    assert client.post('/api/modules/gdi/run', json={'project': '0' * 32}).status_code == 404
    assert client.get('/api/unknown').status_code == 404


def test_missing_dataset_is_409(tmp_path):
    projects = Projects(tmp_path)
    pid = projects.store.create('Пустой')
    r = TestClient(create_app(projects)).post('/api/modules/gdi/run', json={'project': pid, 'params': {}})
    assert r.status_code == 409 and 'ГДИ' in r.json()['error']
