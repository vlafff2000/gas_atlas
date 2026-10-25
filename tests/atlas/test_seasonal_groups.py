"""«Сезонные показатели по группам»: доли и Δ доли, простои, баланс закачки и отбора, подсказки с причиной."""
import io
from pathlib import Path

import pandas as pd
import pytest

from atlas.contract import Data
from atlas.domain import DatasetKind
from atlas.modules import _seasonal as S
from atlas.modules import _seasonal_groups as G
from atlas.modules.seasonal_groups import SeasonalGroupsModule

PROD = DatasetKind.PRODUCTION
GROUPS = {'1': 'ГСП 1', '2': 'ГСП 2', '3': 'ГСП 2'}


def frame():
    """Закачка 2022 и 2023, отбор 2022-2023 (цикл 2022). Скважина 1 — ГСП 1; 2 и 3 — ГСП 2. Скважина 2 простаивает 3 суток отбора."""
    rows = []
    for kind, period, start, rates in (
            ('injection', '2022', '2022-05-01', {'1': 100000, '2': 50000}),
            ('withdrawal', '2022-2023', '2022-11-01', {'1': 200000, '2': 100000}),
            ('injection', '2023', '2023-05-01', {'1': 150000, '2': 50000})):
        for d in pd.date_range(start, periods=10):
            for well, rate in rates.items():
                q = float(rate)
                if kind == 'withdrawal' and well == '2' and d.day <= 3:
                    q = 0.0
                rows.append({'well': well, 'date': d, 'q': q, 'kind': kind, 'period': period})
    return pd.DataFrame(rows)


def seasons(df=None):
    return S.season_table(frame() if df is None else df)


def wells_of(df=None):
    df = frame() if df is None else df
    return G.well_table(df, seasons(df), GROUPS, 50.0)


def test_group_share_and_delta():
    df = frame()
    g = G.group_table(wells_of(df), seasons(df)).set_index(['kind', 'period', 'group'])
    assert g.loc[('injection', '2022', 'ГСП 1'), 'share'] == pytest.approx(100 * 2 / 3)
    assert g.loc[('injection', '2023', 'ГСП 1'), 'share'] == pytest.approx(75.0)
    assert g.loc[('injection', '2023', 'ГСП 1'), 'delta_share'] == pytest.approx(75 - 200 / 3)
    assert pd.isna(g.loc[('injection', '2022', 'ГСП 1'), 'delta_share'])            # первый сезон


def test_group_filter_keeps_shares_of_all_groups():
    df = frame()
    w = wells_of(df)
    only = G.group_table(w, seasons(df), ['ГСП 1']).set_index(['kind', 'period', 'group'])
    assert list(only.index.get_level_values('group').unique()) == ['ГСП 1']
    assert only.loc[('injection', '2022', 'ГСП 1'), 'share'] == pytest.approx(100 * 2 / 3)


def test_idle_wells_and_days():
    df = frame()
    w = wells_of(df).set_index(['kind', 'period', 'well'])
    row = w.loc[('withdrawal', '2022-2023', '2')]
    assert row.idle_days == 3 and row.flow_days == 7
    assert row.status == '' and row.work_pct == pytest.approx(7 / row.span * 100)
    df2 = df.copy()
    df2.loc[(df2.kind == 'injection') & (df2.well == '2') & (df2.period == '2023'), 'q'] = 0.0
    fully = G.well_table(df2, seasons(df2), GROUPS, 50.0).set_index(['kind', 'period', 'well'])
    assert fully.loc[('injection', '2023', '2')].status == 'простаивала весь сезон'
    g = G.group_table(G.well_table(df2, seasons(df2), GROUPS, 50.0), seasons(df2)).set_index(['kind', 'period', 'group'])
    assert g.loc[('injection', '2023', 'ГСП 2'), 'wells_idle'] == 1 and g.loc[('injection', '2023', 'ГСП 2'), 'wells_work'] == 0


def test_missing_rows_are_not_idle():
    df = frame()
    df = df[~((df.kind == 'withdrawal') & (df.well == '2') & (df.date.dt.day <= 5))]     # нули первых 3 суток сменились пропуском
    w = wells_of(df).set_index(['kind', 'period', 'well'])
    assert w.loc[('withdrawal', '2022-2023', '2')].idle_days == 0 and w.loc[('withdrawal', '2022-2023', '2')].flow_days == 5


def test_excluded_points_are_not_counted():
    df = frame()
    df['_excluded'] = (df.well == '2') & (df.kind == 'withdrawal')
    w = wells_of(df).set_index(['kind', 'period', 'well'])
    assert ('withdrawal', '2022-2023', '2') not in w.index


def test_balance_and_hints_with_reason():
    df = frame()
    all_groups = G.group_table(wells_of(df), seasons(df))
    b = G.balance_table(all_groups, None, 10.0).set_index(['cycle', 'group'])
    assert b.loc[(2022, 'ГСП 1'), 'injection'] == pytest.approx(1.0)
    assert b.loc[(2022, 'ГСП 1'), 'withdrawal'] == pytest.approx(2.0)
    assert b.loc[(2022, 'Итого'), 'injection'] == pytest.approx(1.5)
    assert b.loc[(2022, 'Итого'), 'withdrawal'] == pytest.approx(2.7)                  # скважина 2 простояла 3 суток
    assert b.loc[(2022, 'Итого'), 'ratio'] == pytest.approx(180.0)
    assert 'отбор 180 % закачки' in b.loc[(2022, 'Итого'), 'hint']
    # Кривой порог: при 300 % подсказки нет
    assert G.balance_table(all_groups, None, 300.0).set_index(['cycle', 'group']).loc[(2022, 'Итого'), 'hint'] == ''


def test_balance_reports_missing_side():
    df = frame()
    df = df[df.kind.eq('injection')]
    b = G.balance_table(G.group_table(G.well_table(df, seasons(df), GROUPS, 50.0), seasons(df)), None, 10.0)
    assert b.hint.str.contains('в цикле нет отбора').all()


def test_module_run_and_excel():
    from atlas.render import tables_file
    data = Data({PROD: frame()})
    data.mapping = {w: {'group': g} for w, g in GROUPS.items()}
    r = SeasonalGroupsModule().run(data, {'groups': [], 'balance_tol': 10, 'work_low': 50})
    assert {t.id for t in r.tables} == {'seasonal-groups-share', 'seasonal-groups-balance', 'seasonal-groups-wells'}
    assert {c.id for c in r.charts} >= {'seasonal-groups-share-injection', 'seasonal-groups-balance'}
    assert r.summary and not any(n.level == 'warning' for n in r.notes)
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(tables_file(r.tables, 'x')[0]))
    assert book.sheetnames
    one = SeasonalGroupsModule().run(data, {'groups': ['ГСП 1'], 'balance_tol': 10, 'work_low': 50})
    assert set(one.tables[0].frame.group) == {'ГСП 1'}
    empty = SeasonalGroupsModule().run(Data({PROD: frame().iloc[0:0]}), {})
    assert not empty.tables and empty.notes


KRU_HISTORY = Path(__file__).parent / 'data' / 'revision_history_2021_2025.xlsx'


def test_kru_streaks_for_group():
    from atlas.engine.modules import revision_report as rr
    from atlas.modules import _group_problems as P
    dataset = rr.parse_history(KRU_HISTORY)
    kru = P.kru_table(dataset, ['ГСП 1'], 2)
    assert not kru.empty and set(kru.gsp) == {'ГСП-1'}
    assert kru.streak.is_monotonic_decreasing                       # сортировка: самые длинные серии сверху
    assert (kru.hint.ne('') == kru.streak.ge(2)).all()              # подсказка есть ровно там, где серия не короче порога
    top = kru.iloc[0]
    assert top.hint.startswith(f'вскрывалась {top.streak} сезона подряд') or top.streak < 2
    assert P.kru_table(dataset, ['Несуществующая'], 2).empty


def test_module_builds_kru_and_gdi_tables():
    from pathlib import Path as _P
    from atlas.domain import DatasetKind as DK
    from atlas.engine.core.demo import well_demo_frames
    gdi = well_demo_frames()['gdi']
    data = Data({PROD: frame(), DK.GDI: gdi})
    data.mapping = {str(w): {'group': 'ГСП 1'} for w in gdi.well.unique()}
    r = SeasonalGroupsModule().run(data, {'groups': [], 'balance_tol': 10, 'work_low': 50, 'kru_history': str(KRU_HISTORY),
                                          'kru_streak': 2, 'gdi_change': 10})
    ids = {t.id for t in r.tables}
    assert 'seasonal-groups-kru' in ids
    assert 'seasonal-groups-gdi' in ids
    bad = SeasonalGroupsModule().run(data, {'groups': [], 'balance_tol': 10, 'work_low': 50, 'kru_history': str(_P('/нет/файла.xlsx'))})
    assert any(n.level == 'warning' and 'КРУ' in n.text for n in bad.notes)
