"""Excel «Σ hi × Кг» для ГИС: все промежуточные числа — формулы, а не значения.

Три листа: «Интервалы» (hi × Кг по каждому интервалу), «Исследования» (сумма и число интервалов по исследованию и пачке,
SUMIFS/COUNTIFS по листу интервалов) и «Итог» (исследования в строках, пачки в столбцах, сумма по пачкам). Расчёт — в
atlas/engine/modules/gis_hikg.py; здесь только запись формул, поэтому цифры в Excel совпадают с расчётом Атласа.
"""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .engine.core.config import ordered
from .engine.modules import gis_hikg

NO_PACK = '—'
IV, SH = "'Интервалы'", "'Исследования'"    # имена листов в формулах (с кавычками: в них кириллица)
HEAD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')


def _rows(prepared: pd.DataFrame) -> pd.DataFrame:
    """Интервалы выбранных сезонов, как в Атласе: строка на интервал, пустая пачка — «—»."""
    out = prepared.copy()
    out['pack'] = out['pack'].astype(str).replace({'nan': NO_PACK, '': NO_PACK})
    out['investigation'] = out['investigation'].astype(str)
    return out.sort_values(['well', 'date', 'investigation', 'from_m'], kind='stable').reset_index(drop=True)


def _head(ws, names: list[str], widths: list[int] | None = None) -> None:
    ws.append(names)
    for col, name in enumerate(names, 1):
        ws.cell(1, col).font = HEAD
        ws.cell(1, col).alignment = WRAP
        ws.column_dimensions[get_column_letter(col)].width = (widths or [14] * len(names))[col - 1]
    ws.freeze_panes = 'A2'


def build(prepared: pd.DataFrame) -> bytes:
    """Книга: интервалы выбранных сезонов (уже отфильтрованные и подготовленные prepare()). Пустой вход — ValueError."""
    if prepared.empty:
        raise ValueError('За выбранные сезоны нет интервалов ГИС: проверьте сезоны и скважины.')
    rows = _rows(prepared)
    wb = Workbook()

    # «Интервалы»: hi × Кг — формула по каждой строке; Кг 0 и пустой — пустая ячейка
    ws = wb.active
    ws.title = 'Интервалы'
    _head(ws, ['Скважина', 'Сезон', 'Дата', 'Исследование', 'Пачка', 'От, м', 'До, м', 'hi, м', 'Кг', 'hi × Кг'],
          [11, 14, 12, 14, 16, 9, 9, 9, 9, 11])
    for i, r in enumerate(rows.itertuples(index=False), 2):
        ws.append([str(r.well), r.season, r.date.to_pydatetime(), r.investigation, r.pack,
                   _num(r.from_m), _num(r.to_m), _num(r.hi_m), _num(r.kg)])
        ws.cell(i, 3).number_format = 'DD.MM.YYYY'
        ws.cell(i, 10).value = f'=IF(AND(ISNUMBER(H{i}),ISNUMBER(I{i})),IF(I{i}>0,H{i}*I{i},""),"")'
        ws.cell(i, 10).number_format = '0.00'
    last = len(rows) + 1

    # «Исследования»: ключ — скважина, исследование, пачка; значения — SUMIFS по листу интервалов
    keys = rows[['well', 'season', 'date', 'investigation', 'pack']].drop_duplicates().reset_index(drop=True)
    wi = wb.create_sheet('Исследования')
    _head(wi, ['Скважина', 'Сезон', 'Дата', 'Исследование', 'Пачка', 'Интервалов с Кг>0', 'Σ hi × Кг'],
          [11, 14, 12, 14, 16, 16, 14])
    for i, r in enumerate(keys.itertuples(index=False), 2):
        wi.append([str(r.well), r.season, r.date.to_pydatetime(), r.investigation, r.pack])
        wi.cell(i, 3).number_format = 'DD.MM.YYYY'
        crit = f'{IV}!$A$2:$A${last},A{i},{IV}!$D$2:$D${last},D{i},{IV}!$E$2:$E${last},E{i}'
        wi.cell(i, 6).value = f'=COUNTIFS({crit},{IV}!$J$2:$J${last},">0")'
        wi.cell(i, 7).value = f'=SUMIFS({IV}!$J$2:$J${last},{crit})'
        wi.cell(i, 7).number_format = '0.00'
    key_last = len(keys) + 1

    # «Итог»: строка на исследование, столбец на пачку и сумма по пачкам
    items = keys[['well', 'season', 'date', 'investigation']].drop_duplicates().reset_index(drop=True)
    packs = ordered(keys['pack'].unique())
    wt = wb.create_sheet('Итог')
    _head(wt, ['Скважина', 'Сезон', 'Дата', 'Исследование', *packs, 'Всего, Σ hi × Кг'], [11, 14, 12, 14] + [14] * len(packs) + [18])
    for i, r in enumerate(items.itertuples(index=False), 2):
        wt.append([str(r.well), r.season, r.date.to_pydatetime(), r.investigation])
        wt.cell(i, 3).number_format = 'DD.MM.YYYY'
        for j, _ in enumerate(packs, 5):
            col = get_column_letter(j)
            wt.cell(i, j).value = (f'=SUMIFS({SH}!$G$2:$G${key_last},{SH}!$A$2:$A${key_last},A{i},'
                                   f'{SH}!$D$2:$D${key_last},D{i},{SH}!$E$2:$E${key_last},{col}$1)')
            wt.cell(i, j).number_format = '0.00'
        wt.cell(i, 5 + len(packs)).value = f'=SUM(E{i}:{get_column_letter(4 + len(packs))}{i})'
        wt.cell(i, 5 + len(packs)).number_format = '0.00'
        wt.cell(i, 5 + len(packs)).font = HEAD
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(value):
    """Число для ячейки; пустое и нечисловое — пусто, без текста «nan»."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return None if pd.isna(number) else number
