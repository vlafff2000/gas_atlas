"""Ревизия КРУ/КРЭ → многолетняя история перечней скважин и аналитика повторяемости (табл. 1.9 отчёта АН).

Расчётное ядро перенесено из скрипта ``well_report.py`` (РЕД ОС / Windows, отдельная программа) без изменения
правил подсчёта. Убраны: графический мастер (tkinter), графики и Word-отчёт (их делает Атлас), запись промежуточных
файлов. Чтение Word — стандартной библиотекой (``zipfile`` + ``xml``), без python-docx: его зависимость lxml не
собирается под glibc 2.17 (РЕД ОС 7.3). Правила входа — README_RU.md исходного скрипта.

Вход: история — книга .xlsx (ГСП | количество | номера | … по сезонам); ревизия нового сезона — .docx с таблицами
«№ скв.» (или .xlsx). Выход: ``Dataset`` с ``analyse`` и ``export_excel``.
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, date
from pathlib import Path
from typing import Any
import math
import re
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class InputError(ValueError):
    """Ошибка входных данных, которую необходимо исправить в Excel/настройках."""


class Cancelled(Exception):
    pass


SEASON_RE = re.compile(r'(?<!\d)((?:19|20|21)\d{2})\s*[/–—-]\s*(\d{4}|\d{2})(?!\d)')
GSP_RE = re.compile(r'^\s*Г\s*С\s*П\s*[-–—№:]?\s*(\d+)\s*$', re.I)
UNKNOWN = {'', 'н/д', 'нет данных', 'нд', 'n/a', 'na'}
EMPTY = {'-', '–', '—', 'нет', 'отсутствуют'}
WELL_RE = re.compile(r'[0-9]+[А-ЯA-Z]*(?:[/\-][0-9]+[А-ЯA-Z]*)*', re.I)


def text(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace('\xa0', ' ').strip()


def sort_key(value: str):
    return tuple((0, int(p)) if p.isdigit() else (1, p)
                 for p in re.split(r'(\d+)', str(value)) if p)


def normalize_season(value: Any) -> str:
    match = SEASON_RE.search(text(value))
    if not match:
        raise InputError(f'Не распознан сезон «{value}». Пример: 2025/26 или 2025-2026.')
    first = int(match[1])
    second = int(match[2])
    if len(match[2]) == 2:
        second += first // 100 * 100
        if second < first:
            second += 100
    if second != first + 1:
        raise InputError(f'Сезон должен состоять из соседних лет: «{value}».')
    return f'{first}/{str(second)[-2:]}'


def seasons_in(value: Any) -> list[str]:
    found = []
    for match in SEASON_RE.finditer(text(value)):
        try:
            season = normalize_season(match[0])
        except InputError:
            continue
        if season not in found:
            found.append(season)
    return found


def gsp_name(value: Any) -> str | None:
    match = GSP_RE.fullmatch(text(value))
    return f'ГСП-{int(match[1])}' if match else None


def parse_wells(value: Any, location: str = '') -> list[str]:
    """Не извлекает произвольные цифры из текста и не делит составной номер 54/80."""
    if isinstance(value, (datetime, date)):
        raise InputError(f'{location}: номер превращён Excel в дату. Задайте текстовый формат ячейки.')
    s = text(value).upper()
    if s.lower() in UNKNOWN | EMPTY or s == '0':
        return []
    s = re.sub(r'\bСКВ(?:АЖИН[АЫ]?)?\.?\s*', '', s)
    s = re.sub(r'\s*/\s*', '/', s)
    s = re.sub(r'\s*[-–—]\s*', '-', s)
    s = s.replace('№', '').strip()
    if re.fullmatch(r'\d+\.0+', s):
        s = s.split('.')[0]
    parts = [p for p in re.split(r'[,;\s]+', s) if p]
    invalid = [p for p in parts if not WELL_RE.fullmatch(p) or p == '0']
    if invalid:
        raise InputError(f'{location}: непонятные номера: {", ".join(invalid[:8])}. '
                         'Разделяйте номера запятыми; диапазоны не разворачиваются.')
    return parts


@dataclass
class Entry:
    raw_count: Any = None
    raw_wells: Any = None
    wells: list[str] = field(default_factory=list)
    known: bool = False
    source: str = ''

    @property
    def unique(self) -> set[str]:
        return set(self.wells)


def is_history_note(value: Any, count_column: bool = False) -> bool:
    """Отделяет поясняющий текст от числа/списка, не требуя точной фразы."""
    if not isinstance(value, str) or text(value).lower() in UNKNOWN | EMPTY:
        return False
    try:
        if count_column:
            float(text(value).replace(',', '.'))
        else:
            parse_wells(value)
        return False
    except (InputError, ValueError):
        # Числовые опечатки вроде «3?» остаются ошибками. Фразы с буквами — пометки.
        return bool(re.search(r'[А-Яа-яЁёA-Za-z]{2,}', value))


def merged_history_note(entry: Entry) -> str | None:
    """Одна пометка может занимать обе графы сезона в Word и Excel."""
    if entry.known:
        return None
    count, wells = text(entry.raw_count), text(entry.raw_wells)
    if is_history_note(count, True) and (wells.lower() in UNKNOWN | EMPTY or count == wells):
        return count
    if is_history_note(wells) and count.lower() in UNKNOWN | EMPTY:
        return wells
    return None


def entry_display_values(entry: Entry) -> tuple[Any, Any]:
    count = entry.raw_count if text(entry.raw_count) else (0 if entry.known else 'н/д')
    wells = entry.raw_wells if text(entry.raw_wells) else ('—' if entry.known else 'нет данных')
    return count, wells


@dataclass
class Dataset:
    seasons: list[str]
    gsps: list[str]
    entries: dict[tuple[str, str], Entry]
    issues: list[str] = field(default_factory=list)
    sources: list[str] = field(default_factory=list)

    def get(self, gsp: str, season: str) -> Entry:
        return self.entries.get((gsp, season), Entry())


def load_sheet(path: Path | str, sheet: str | None = None):
    """Читает вычисленные значения формул; не пытается вычислять формулы сам."""
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise InputError('Нужен файл .xlsx. Сначала сохраните .xls или .csv в формате .xlsx.')
    if not path.is_file():
        raise InputError(f'Файл не найден: {path}')
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise InputError(f'Не удалось открыть {path.name}: {exc}') from exc
    if sheet is None:
        if len(wb.sheetnames) != 1:
            names = ', '.join(wb.sheetnames)
            wb.close()
            raise InputError(f'Выберите лист файла {path.name}: {names}')
        sheet = wb.sheetnames[0]
    if sheet not in wb.sheetnames:
        wb.close()
        raise InputError(f'Нет листа «{sheet}» в {path.name}.')
    return wb, wb[sheet]


def value_at(ws, row: int, col: int):
    """Возвращает якорь объединённой ячейки (важно для заголовков и ГСП)."""
    value = ws.cell(row, col).value
    if value is not None:
        return value
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            return ws.cell(merged.min_row, merged.min_col).value
    return None


def cached_value_guard(path: Path, sheet: str, cells: list[tuple[int, int]]):
    # Отсутствующий кэш формулы не должен превратиться в «нет данных».
    if not cells:
        return
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        ws = wb[sheet]
        formulas = [ws.cell(r, c).coordinate for r, c in cells
                    if ws.cell(r, c).data_type == 'f']
        if formulas:
            raise InputError('Нет сохранённых результатов формул в ячейках '
                             f'{", ".join(formulas[:12])}. Откройте файл в Excel/LibreOffice, '
                             'пересчитайте и сохраните его либо вставьте значения.')
    finally:
        wb.close()


def history_layout(ws) -> dict:
    gsp_cells = [(c.row, c.column) for row in ws.iter_rows() for c in row if gsp_name(c.value)]
    if not gsp_cells:
        raise InputError('В сводной не найдены строки «ГСП-1», «ГСП-2» и т. д.')
    gcol = Counter(c for _, c in gsp_cells).most_common(1)[0][0]
    first = min(r for r, c in gsp_cells if c == gcol)
    candidates: dict[str, set[int]] = {}
    for row in ws.iter_rows(max_row=first - 1):
        for cell in row:
            if cell.column <= gcol:
                continue
            for season in seasons_in(cell.value):
                candidates.setdefault(season, set()).add(cell.column)
    if not candidates:
        raise InputError('Не найдены заголовки сезонов. Укажите столбцы вручную в настройках импорта.')
    pairs = []
    for season, cols in candidates.items():
        start = min(cols)
        if len(cols) > 2 or (len(cols) == 2 and max(cols) != start + 1):
            raise InputError(f'Неоднозначные столбцы сезона {season}. Задайте сопоставление вручную.')
        if start + 1 > ws.max_column:
            raise InputError(f'Для сезона {season} нет пары столбцов «Количество» и «Номера».')
        headers = ' '.join(text(ws.cell(r, start).value).lower() for r in range(1, first))
        right = ' '.join(text(ws.cell(r, start + 1).value).lower() for r in range(1, first))
        count_col, wells_col = start, start + 1
        if ('номер' in headers or '№' in headers or 'no ' in headers) and 'кол' in right:
            count_col, wells_col = start + 1, start
        pairs.append({'season': season, 'count_col': count_col, 'wells_col': wells_col})
    used = [p[k] for p in pairs for k in ('count_col', 'wells_col')]
    if len(used) != len(set(used)):
        raise InputError('Пары столбцов сезонов пересекаются. Укажите их вручную.')
    return {'gsp_col': gcol, 'first_row': first, 'last_row': ws.max_row, 'pairs': pairs}


def check_entry(entry: Entry, issues: list[str]):
    if not entry.known:
        return
    duplicates = [f'{w} × {n}' for w, n in Counter(entry.wells).items() if n > 1]
    if duplicates:
        issues.append(f'{entry.source}: повторы в исходном списке: {", ".join(duplicates)}. '
                      'В аналитике каждый номер учтён один раз за сезон.')
    if entry.raw_count is not None:
        try:
            n = float(str(entry.raw_count).replace(',', '.'))
            if not math.isfinite(n) or not n.is_integer() or n < 0:
                raise ValueError
        except (ValueError, TypeError):
            raise InputError(f'{entry.source}: количество должно быть целым неотрицательным числом.')
        if int(n) != len(entry.unique):
            issues.append(f'{entry.source}: в исходной графе количества {int(n)}, '
                          f'уникальных номеров {len(entry.unique)} (записей {len(entry.wells)}). '
                          'Историческое значение сохранено, аналитика рассчитана по номерам.')


def parse_history(path: Path | str, sheet: str | None = None, layout: dict | None = None) -> Dataset:
    path = Path(path)
    wb, ws = load_sheet(path, sheet)
    try:
        layout = layout or history_layout(ws)
        pairs = layout['pairs']
        season_list = [normalize_season(p['season']) for p in pairs]
        if len(season_list) != len(set(season_list)):
            raise InputError('Сезон повторяется в настройках исторической таблицы.')
        entries, issues, gsps, blanks = {}, [], [], []
        for r in range(layout['first_row'], layout.get('last_row', ws.max_row) + 1):
            gsp = gsp_name(ws.cell(r, layout['gsp_col']).value)
            if not gsp:
                continue
            if gsp in gsps:
                raise InputError(f'{ws.title}, строка {r}: повторная строка {gsp} в сводной.')
            gsps.append(gsp)
            for p, season in zip(pairs, season_list):
                raw_count = ws.cell(r, p['count_col']).value
                raw_wells = ws.cell(r, p['wells_col']).value
                # Вертикально/горизонтально объединённая пометка относится ко всем покрытым ячейкам.
                if raw_count is None and is_history_note(value_at(ws, r, p['count_col']), True):
                    raw_count = value_at(ws, r, p['count_col'])
                if raw_wells is None and is_history_note(value_at(ws, r, p['wells_col'])):
                    raw_wells = value_at(ws, r, p['wells_col'])
                loc = f'{gsp}, {season} ({ws.title}!{get_column_letter(p["wells_col"])}{r})'
                for col in (p['count_col'], p['wells_col']):
                    if ws.cell(r, col).value is None:
                        blanks.append((r, col))
                if is_history_note(raw_count, True) or is_history_note(raw_wells):
                    notes = list(dict.fromkeys(text(v) for v in (raw_count, raw_wells) if text(v)))
                    issues.append(f'{loc}: текстовая пометка «{"; ".join(notes)}». '
                                  'Исходный текст сохранён; сезон для этого ГСП исключён из числовой аналитики.')
                    entries[(gsp, season)] = Entry(raw_count, raw_wells, [], False, loc)
                    continue
                wells = parse_wells(raw_wells, loc)
                count_s = text(raw_count).lower()
                list_s = text(raw_wells).lower()
                count_missing = count_s in UNKNOWN | EMPTY
                if not wells and not count_missing and count_s not in {'0', '0.0'}:
                    raise InputError(f'{loc}: указано количество {raw_count}, но список номеров пуст.')
                known = bool(wells) or (count_s in {'0', '0.0'} and list_s not in UNKNOWN - {''})
                count = raw_count if not count_missing else (len(set(wells)) if wells else None)
                if wells and count_missing:
                    issues.append(f'{loc}: количество отсутствует; рассчитано по уникальным номерам: {count}.')
                entry = Entry(count, raw_wells, wells, known, loc)
                check_entry(entry, issues)
                entries[(gsp, season)] = entry
        if not gsps:
            raise InputError('В выбранном диапазоне истории нет строк ГСП.')
        cached_value_guard(path, ws.title, blanks)
        return Dataset(sorted(season_list), sorted(gsps, key=sort_key), entries, issues,
                       [f'История: {path.name}, лист «{ws.title}», строки {layout["first_row"]}–'
                        f'{layout.get("last_row", ws.max_row)}.'])
    finally:
        wb.close()


def revision_layout(ws) -> dict:
    candidates = []
    word_rows = getattr(ws, '_docx_row_locations', None)
    limit = ws.max_row if word_rows is not None else min(ws.max_row, 100)
    for row in ws.iter_rows(max_row=limit):
        if word_rows is not None and row[0].row not in word_rows:
            continue
        for cell in row:
            s = text(cell.value).lower()
            if 'скв' in s and 'кол' not in s and len(s) < 65 and ('№' in s or 'номер' in s or 'no' in s):
                candidates.append((cell.row, cell.column))
    cols = {c for _, c in candidates}
    if len(cols) != 1:
        raise InputError('Не удалось однозначно определить столбец «№ скв.». Укажите его вручную.')
    r, col = min(candidates)
    return {'well_col': col, 'first_row': r + 1, 'last_row': ws.max_row, 'gsp_col': None}


def parse_table_numbers(value: Any) -> list[int] | None:
    if value is None or value == '' or value == []:
        return None
    try:
        parts = re.split(r'[,;\s]+', str(value).strip()) if isinstance(value, str) else value
        numbers = [int(p) for p in parts if str(p).strip()]
    except (TypeError, ValueError) as exc:
        raise InputError('Номера таблиц Word задаются целыми числами через запятую: 1,2,3.') from exc
    if not numbers or any(n < 1 for n in numbers) or len(numbers) != len(set(numbers)):
        raise InputError('Номера таблиц Word должны быть положительными и не повторяться.')
    return sorted(numbers)


W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _docx_text(element) -> str:
    """Текст абзаца или ячейки как в python-docx: табуляция и перенос строки сохраняются."""
    parts = []
    for node in element.iter():
        if node.tag == W_NS + 't':
            parts.append(node.text or '')
        elif node.tag == W_NS + 'tab':
            parts.append('\t')
        elif node.tag in (W_NS + 'br', W_NS + 'cr'):
            parts.append('\n')
    return ''.join(parts)


def _docx_row(tr, anchors: dict[int, str]) -> list[str]:
    """Значения строки по сетке столбцов: объединённая по горизонтали ячейка даёт текст в первом столбце, пусто — в остальных;
    продолжение объединения по вертикали повторяет текст верхней ячейки (как ``row.cells`` в python-docx)."""
    values = []
    grid_before = tr.find(f'{W_NS}trPr/{W_NS}gridBefore')
    values += [''] * int(grid_before.get(W_NS + 'val', '0')) if grid_before is not None else []
    for tc in tr.findall(W_NS + 'tc'):
        props = tc.find(W_NS + 'tcPr')
        span_el = props.find(W_NS + 'gridSpan') if props is not None else None
        span = int(span_el.get(W_NS + 'val', '1')) if span_el is not None else 1
        merge = props.find(W_NS + 'vMerge') if props is not None else None
        start = len(values)
        if merge is not None and merge.get(W_NS + 'val', 'continue') == 'continue':
            text = anchors.get(start, '')
        else:
            text = '\n'.join(_docx_text(p).strip() for p in tc.findall(W_NS + 'p')).strip()
            for col in range(start, start + span):
                anchors[col] = text
        values.append(text)
        values += [''] * (span - 1)
    return values


def docx_body(path: Path):
    """Абзацы и таблицы Word в порядке тела документа: ('p', текст) или ('tbl', строки)."""
    try:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read('word/document.xml'))
    except (zipfile.BadZipFile, KeyError, ET.ParseError, OSError) as exc:
        raise InputError(f'Не удалось прочитать Word-файл {path.name}: {exc}') from exc
    body = root.find(W_NS + 'body')
    if body is None:
        raise InputError(f'Не удалось прочитать Word-файл {path.name}: нет тела документа.')
    for element in list(body):
        if element.tag == W_NS + 'p':
            yield 'p', _docx_text(element).strip()
        elif element.tag == W_NS + 'tbl':
            anchors: dict[int, str] = {}
            yield 'tbl', [_docx_row(tr, anchors) for tr in element.findall(W_NS + 'tr')]


def docx_revision_book(path: Path | str, tables: list[int] | None = None):
    """Читает редактируемые Word-таблицы в память для общего предпросмотра/парсера.

    Это адаптер, а не промежуточный файл Excel. Порядок абзацев и таблиц,
    метки ГСП над таблицей, объединённые ячейки и адреса исходных строк сохраняются.
    Таблицы без повторной шапки можно включить явным списком номеров.
    Чтение — стандартной библиотекой (без python-docx и lxml: нет сборок под glibc 2.17 для РЕД ОС 7.3).
    """
    path = Path(path)
    if not path.is_file():
        raise InputError(f'Файл не найден: {path}')
    records, pending, current_gsp = [], [], None
    for kind, value in docx_body(path):
        if kind == 'p':
            if value:
                pending.append(value)
                current_gsp = gsp_name(value) or current_gsp
            continue
        rows = value
        cells = [v.lower() for row in rows[:8] for v in row]
        has_well_header = any('скв' in s and 'кол' not in s and len(s) < 65
                              and ('№' in s or 'номер' in s or 'no' in s) for s in cells)
        has_description = any(re.search(r'разруш|износ|поврежд|дефект|ревизи', s) for s in cells)
        records.append({'number': len(records) + 1, 'rows': rows, 'prefix': pending,
                        'gsp': current_gsp, 'header': has_well_header,
                        'strong': has_well_header and has_description})
        pending = []
        for row in rows:
            for v in row:
                current_gsp = gsp_name(v) or current_gsp
    if not records:
        raise InputError('В .docx не найдены редактируемые таблицы. Изображения/сканы таблиц не распознаются.')
    selected = parse_table_numbers(tables)
    if selected is None:
        selected = ([r['number'] for r in records if r['strong']]
                    or [r['number'] for r in records if r['header']]
                    or [r['number'] for r in records])
    if max(selected) > len(records):
        raise InputError(f'В Word-файле только {len(records)} таблиц. Проверьте выбранные номера.')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ревизия Word'
    ws._docx_row_locations = {}
    for record in records:
        if record['number'] not in selected:
            continue
        if record['gsp']:
            ws.append([record['gsp']])
        for paragraph in record['prefix']:
            ws.append([paragraph])
        for i, values in enumerate(record['rows'], 1):
            ws.append(values or [''])
            ws._docx_row_locations[ws.max_row] = f'таблица Word {record["number"]}, строка {i}'
    wb._docx_selected = selected
    wb._docx_table_count = len(records)
    # Уведомляем о таблицах без шапки: пользователь может включить продолжения вручную.
    wb._docx_unselected = [r['number'] for r in records if r['number'] not in selected]
    return wb


def parse_revision(path: Path | str, sheet: str | None = None, layout: dict | None = None,
                   filter_regex: str = '') -> tuple[dict[str, Entry], list[str], str, list[str]]:
    path = Path(path)
    is_word = path.suffix.lower() == '.docx'
    if is_word:
        wb = docx_revision_book(path, (layout or {}).get('docx_tables'))
        ws = wb.active
    else:
        wb, ws = load_sheet(path, sheet)
    try:
        if not layout or not {'well_col', 'first_row'} <= layout.keys():
            layout = {**revision_layout(ws), **(layout or {})}
        pattern = re.compile(filter_regex, re.I) if filter_regex else None
        entries: dict[str, Entry] = {}
        issues, blanks, hints = [], [], []
        if is_word and wb._docx_unselected:
            issues.append('Word: не включены таблицы ' + ', '.join(map(str, wb._docx_unselected))
                          + '. Если это продолжения ревизии, укажите их номера в настройках импорта.')
        current = None
        valid_rows = 0
        last = layout.get('last_row', ws.max_row)
        for row in ws.iter_rows(max_row=min(layout['first_row'] + 2, ws.max_row)):
            for cell in row:
                for season in seasons_in(cell.value):
                    if season not in hints:
                        hints.append(season)
        # Читаем заголовки ГСП до первой строки данных, чтобы не потерять контекст секции.
        for r in range(1, last + 1):
            gcol = layout.get('gsp_col')
            if gcol:
                groups = {gsp_name(value_at(ws, r, gcol))} - {None}
            else:
                groups = {gsp_name(c.value) for c in ws[r]} - {None}
            if len(groups) > 1:
                raise InputError(f'{ws.title}, строка {r}: несколько ГСП в одной строке.')
            if groups:
                current = groups.pop()
            if r < layout['first_row']:
                continue
            if current is not None:
                entries.setdefault(current, Entry(None, None, [], False,
                                                   f'{current} ({ws.title}, ревизия)'))
            raw = ws.cell(r, layout['well_col']).value
            if raw is None:
                blanks.append((r, layout['well_col']))
                continue
            s = text(raw)
            if gsp_name(raw) or ('скв' in s.lower() and ('№' in s or 'номер' in s.lower())):
                continue
            if s.lower() in UNKNOWN | EMPTY or s == '0':
                continue
            if re.match(r'^(итого|всего)\b', s, re.I):
                continue
            loc = (f'{ws._docx_row_locations.get(r, "строка предпросмотра " + str(r))}, '
                   f'столбец {get_column_letter(layout["well_col"])}' if is_word
                   else f'{ws.title}!{get_column_letter(layout["well_col"])}{r}')
            wells = parse_wells(raw, loc)
            if len(wells) != 1:
                raise InputError(f'{loc}: ожидается один номер скважины в строке ревизии.')
            if current is None:
                raise InputError(f'{loc}: скважина {wells[0]} расположена до заголовка ГСП.')
            valid_rows += 1
            entries[current].known = True
            if pattern:
                row_text = ' '.join(text(c.value) for c in ws[r] if c.column != layout['well_col'])
                if not pattern.search(row_text):
                    continue
            entries[current].wells.extend(wells)
        if not is_word:
            cached_value_guard(path, ws.title, blanks)
        if not entries or not valid_rows:
            raise InputError('В ревизии не найдено скважин. Проверьте лист, строки, столбцы и фильтр.')
        for entry in entries.values():
            if not entry.known:
                issues.append(f'{entry.source}: раздел есть, но строк со скважинами нет; отмечено «нет данных».')
                continue
            entry.raw_count = None
            check_entry(entry, issues)
            entry.wells = list(dict.fromkeys(entry.wells))
            entry.raw_count = len(entry.wells)
            entry.raw_wells = ', '.join(entry.wells)
        source_place = ('таблицы Word ' + ', '.join(map(str, wb._docx_selected))
                        + ', строки предпросмотра' if is_word else f'лист «{ws.title}», строки')
        source = (f'Новый сезон: {path.name}, {source_place} {layout["first_row"]}–{last}. '
                  + (f'Фильтр текста строк: {filter_regex}.' if filter_regex else 'Включены все строки с номерами скважин.'))
        return entries, issues, source, hints
    finally:
        wb.close()


def append_season(history: Dataset, season: str, revision: dict[str, Entry],
                  issues: list[str], source: str, absent_zero: bool = False) -> Dataset:
    season = normalize_season(season)
    if season in history.seasons:
        raise InputError(f'Сезон {season} уже присутствует в истории. Выберите таблицу до этого сезона.')
    if history.seasons and season < max(history.seasons):
        raise InputError('Новый сезон должен быть позднее последнего исторического сезона.')
    entries = dict(history.entries)
    gsps = sorted(set(history.gsps) | set(revision), key=sort_key)
    combined_issues = list(history.issues) + list(issues)
    for gsp in gsps:
        if gsp in revision:
            entries[(gsp, season)] = revision[gsp]
        elif absent_zero:
            entries[(gsp, season)] = Entry(0, '—', [], True, 'Ноль по выбранному правилу отсутствия ГСП')
            combined_issues.append(f'{gsp}, {season}: ГСП отсутствует в ревизии; принят ноль по выбранному правилу.')
        else:
            entries[(gsp, season)] = Entry(source='ГСП отсутствует в ревизии')
            combined_issues.append(f'{gsp}, {season}: ГСП отсутствует в ревизии; отмечено «нет данных».')
    seasons = sorted(history.seasons + [season])
    for a, b in zip(seasons, seasons[1:]):
        if int(b[:4]) != int(a[:4]) + 1:
            combined_issues.append(f'Между {a} и {b} есть пропуск календарных сезонов. '
                                   'Повторяемость год к году для этого перехода не рассчитывается.')
    well_gsps: dict[str, set[str]] = {}
    for (g, _), e in entries.items():
        for w in e.unique:
            well_gsps.setdefault(w, set()).add(g)
    for w, groups in well_gsps.items():
        if len(groups) > 1:
            combined_issues.append(f'Номер {w} встречается в нескольких ГСП: {", ".join(sorted(groups))}. '
                                   'В общем итоге это разные пары «ГСП + номер».')
    return Dataset(seasons, gsps, entries, combined_issues, history.sources + [source])


def consecutive(a: str, b: str) -> bool:
    return int(b[:4]) == int(a[:4]) + 1


def longest_streak(present: list[bool], seasons: list[str]) -> int:
    best = run = 0
    for i, yes in enumerate(present):
        run = (run + 1 if i and consecutive(seasons[i - 1], seasons[i]) else 1) if yes else 0
        best = max(best, run)
    return best


def analyse(data: Dataset) -> dict:
    """Единица учёта во всех метриках — пара ГСП + нормализованный номер."""
    per_well, summary, transitions = [], [], []
    for g in data.gsps:
        sets = [data.get(g, s).unique for s in data.seasons]
        known = [data.get(g, s).known for s in data.seasons]
        union = set().union(*sets)
        n_known = sum(known)
        for w in sorted(union, key=sort_key):
            flags = [w in present for present in sets]
            indices = [i for i, yes in enumerate(flags) if yes]
            k = len(indices)
            streak = longest_streak(flags, data.seasons)
            per_well.append({'gsp': g, 'well': w, 'k': k, 'n': n_known,
                             'share': k / n_known if n_known else None,
                             'all': n_known == len(data.seasons) and k == n_known,
                             'first': data.seasons[indices[0]], 'last': data.seasons[indices[-1]],
                             'streak': streak, 'flags': [int(f) if ok else None for f, ok in zip(flags, known)]})
        rows = [r for r in per_well if r['gsp'] == g]
        events = sum(len(s) for s in sets)
        summary.append({'gsp': g, 'coverage': n_known, 'unique': len(union),
                        'events': events, 'mean': events / n_known if n_known else None,
                        'all': sum(r['all'] for r in rows) if n_known == len(data.seasons) else None,
                        'once': sum(r['k'] == 1 for r in rows),
                        'repeated': sum(r['k'] >= 2 for r in rows),
                        'repeated_share': sum(r['k'] >= 2 for r in rows) / len(union) if union else None})
        for i in range(1, len(data.seasons)):
            prev, curr = data.seasons[i - 1:i + 1]
            t = {'gsp': g, 'from': prev, 'to': curr, 'available': False}
            if known[i - 1] and known[i] and consecutive(prev, curr):
                a, b = sets[i - 1:i + 1]
                earlier = set().union(*sets[:i])
                entered = b - a
                t.update(available=True, retained=len(a & b), entered=len(entered),
                         new=len(b - earlier), returned=len(entered & earlier),
                         absent=len(a - b), previous=len(a), current=len(b),
                         change=len(b) - len(a),
                         change_pct=(len(b) - len(a)) / len(a) if a else None,
                         retention=len(a & b) / len(a) if a else None,
                         jaccard=len(a & b) / len(a | b) if a | b else None,
                         retained_wells=sorted(a & b, key=sort_key),
                         new_wells=sorted(b - earlier, key=sort_key),
                         returned_wells=sorted(entered & earlier, key=sort_key),
                         absent_wells=sorted(a - b, key=sort_key))
            transitions.append(t)
    totals = []
    for s in data.seasons:
        known = sum(data.get(g, s).known for g in data.gsps)
        totals.append({'season': s, 'count': sum(len(data.get(g, s).unique) for g in data.gsps),
                       'complete': known == len(data.gsps), 'known': known})
    return {'wells': per_well, 'summary': summary, 'transitions': transitions, 'totals': totals}


def percent(value: float | None) -> str:
    return 'н/д' if value is None else f'{100 * value:.1f}%'.replace('.', ',')


def number(value: Any) -> str:
    if value is None:
        return 'н/д'
    return f'{value:.1f}'.replace('.', ',') if isinstance(value, float) else str(value)


def excel_style(ws, header_rows: int = 1, widths: list[float] | None = None):
    side = Side(style='thin', color='D9D9D9')
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name='Calibri', size=11, bold=c.row <= header_rows,
                          color='FFFFFF' if c.row <= header_rows else '20252B')
            c.fill = PatternFill('solid', fgColor='203D54' if c.row <= header_rows
                                else ('F2F5F7' if c.row % 2 == 0 else 'FFFFFF'))
            c.alignment = Alignment(horizontal='left' if c.column == 1 else 'center',
                                    vertical='center', wrap_text=True)
            c.border = Border(left=side, right=side, top=side, bottom=side)
        ws.row_dimensions[row[0].row].height = 34 if row[0].row <= header_rows else 30
    if widths:
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = f'B{header_rows + 1}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f'1:{header_rows}'


def export_excel(data: Dataset, analytics: dict, path: Path):
    """Новая переносимая сводная. Исходные книги не изменяются."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Сводная'
    ws.cell(1, 1, 'Номер ГСП')
    ws.merge_cells(start_row=1, end_row=2, start_column=1, end_column=1)
    for i, s in enumerate(data.seasons):
        c = 2 + i * 2
        ws.cell(1, c, f'Отбор {s} гг.')
        ws.merge_cells(start_row=1, end_row=1, start_column=c, end_column=c + 1)
        ws.cell(2, c, 'Кол. скв., шт.')
        ws.cell(2, c + 1, '№ скважин')
    for r, g in enumerate(data.gsps, 3):
        ws.cell(r, 1, g)
        for i, s in enumerate(data.seasons):
            e = data.get(g, s)
            c = 2 + 2 * i
            note = merged_history_note(e)
            if note:
                ws.cell(r, c, note)
                ws.merge_cells(start_row=r, end_row=r, start_column=c, end_column=c + 1)
            else:
                count, wells = entry_display_values(e)
                ws.cell(r, c, count)
                ws.cell(r, c + 1, wells)
    excel_style(ws, 2, [15] + [12, 32] * len(data.seasons))
    for r in range(3, ws.max_row + 1):
        lengths = []
        for i, s in enumerate(data.seasons):
            e = data.get(data.gsps[r - 3], s)
            note = merged_history_note(e)
            count, wells = entry_display_values(e)
            lengths.append(math.ceil(len(note) / 39) if note else
                           max(math.ceil(len(text(count)) / 10), math.ceil(len(text(wells)) / 29)))
        lines = max(lengths)
        ws.row_dimensions[r].height = max(36, 16 * (lines + 1))
        for c in range(3, ws.max_column + 1, 2):
            ws.cell(r, c).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    freq = wb.create_sheet('Частота')
    freq.append(['ГСП', 'Скважина', 'Появлений', 'Сезонов с данными', 'Частота',
                 'Первое появление', 'Последнее появление', 'Макс. серия'] + data.seasons)
    for w in analytics['wells']:
        freq.append([w['gsp'], w['well'], w['k'], w['n'], w['share'], w['first'], w['last'], w['streak']]
                    + [('н/д' if x is None else x) for x in w['flags']])
    excel_style(freq, widths=[15, 14, 13, 18, 13, 19, 19, 14] + [13] * len(data.seasons))
    for row in freq.iter_rows(min_row=2):
        row[4].number_format = '0.0%'
    freq.auto_filter.ref = freq.dimensions
    tr = wb.create_sheet('Переходы')
    tr.append(['ГСП', 'Предыдущий', 'Текущий', 'Повторно', 'Впервые в данных', 'Вернулись',
               'Нет в текущем', 'Изменение', 'Изменение, %', 'Повторяемость', 'Жаккар'])
    for t in analytics['transitions']:
        keys = ['retained', 'new', 'returned', 'absent', 'change', 'change_pct', 'retention', 'jaccard']
        tr.append([t['gsp'], t['from'], t['to']] + [t.get(k, 'н/д') for k in keys])
    excel_style(tr, widths=[14, 16, 16, 14, 18, 14, 18, 15, 18, 18, 14])
    for row in tr.iter_rows(min_row=2):
        for c in row[8:11]:
            c.number_format = '0.0%'
    notes = wb.create_sheet('Примечания')
    notes.append(['Описание'])
    for line in methodology(data) + data.sources + data.issues:
        notes.append([line])
    excel_style(notes, widths=[125])
    for r in range(2, notes.max_row + 1):
        notes.row_dimensions[r].height = max(30, 16 * math.ceil(len(text(notes.cell(r, 1).value)) / 115))
    wb.save(path)
    wb.close()


def methodology(data: Dataset) -> list[str]:
    return [
        'Исторические количества и списки перенесены из исходной таблицы. '
        'Если количество отсутствовало при наличии списка, оно восстановлено по уникальным номерам.',
        'Текстовые пометки в графах сезона, например «Ревизия не проводилась» или «Вскрытие КРУ не выполнялось», '
        'сохраняются в таблице. Для соответствующего ГСП такой сезон считается отсутствием данных, '
        'исключается из знаменателя частоты и не участвует в числовом сравнении соседних сезонов.',
        'Аналитика использует уникальные пары «ГСП + номер скважины» за каждый сезон. '
        'Повтор внутри одного списка не увеличивает аналитическое количество. '
        'Составной номер, например 54/80, считается одним идентификатором; диапазоны не разворачиваются.',
        'Частота равна числу сезонов с появлением скважины, делённому на число сезонов с данными для её ГСП. '
        '«Во всех сезонах» требует полного покрытия всех рассматриваемых сезонов. '
        '«Однократно» означает одно появление среди доступных данных.',
        'Повторяемость между соседними календарными сезонами равна числу общих скважин, '
        'делённому на количество скважин предыдущего сезона. '
        'Сходство Жаккара равно числу общих скважин, делённому на размер объединения двух списков. '
        'При нулевом знаменателе доля не рассчитывается.',
        '«Впервые» означает первое появление в загруженной истории. '
        '«Вернулись» — номера из более ранних доступных сезонов, которых нет в предыдущем сезоне. '
        '«Нет в текущем» означает отсутствие записи в перечне; это не подтверждает устранение причины износа.',
        'Пустые ячейки и отсутствующие ГСП обозначают отсутствие данных. '
        'Нулевое количество с пустым списком означает подтверждённый пустой перечень. '
        'Пустой раздел ревизии также означает отсутствие данных. Если строки ревизии есть, '
        'но ни одна не проходит заданный фильтр, результат равен нулю. '
        'Пропуски разрывают последовательную серию. Графики не соединяют пропущенные календарные сезоны.',
        'Сумма появлений измеряется в «скважина × сезон». Она отличается от количества уникальных скважин. '
        'Доли фонда скважин не рассчитываются: численность всего фонда не задана.'
    ]
