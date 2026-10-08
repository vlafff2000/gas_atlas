"""Local static exports; no Chrome/CDN required. Plotly traces → Matplotlib."""
import base64
import io
import json
import re
import zipfile
import threading
import textwrap
import datetime
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

LOCK=threading.RLock()

def safe_name(name):
    return re.sub(r'[^\w .-]','_',str(name)).strip(' .')[:100] or 'chart'

def safe_table(df):
    d=df.copy()
    for col in d.select_dtypes(include=['object','string']):
        d[col]=d[col].map(lambda v:"'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v)
    return d

def csv_bytes(df):
    return safe_table(df).to_csv(index=False,sep=';',decimal=',').encode('utf-8-sig')

def xlsx_bytes(tables):
    wb=Workbook(write_only=True)
    from openpyxl.cell import WriteOnlyCell
    for i,(name,df) in enumerate(tables.items()):
        name=re.sub(r'[\[\]:*?/\\]','_',name)[:25]+'_'+str(i+1)
        ws=wb.create_sheet(name); ws.freeze_panes='A2'
        headers=[]
        for col in df:
            c=WriteOnlyCell(ws,str(col)); c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='243247'); headers.append(c)
        ws.append(headers)
        for row in safe_table(df).itertuples(index=False,name=None):
            vals=[]
            for v in row:
                if pd.isna(v): v=None
                elif isinstance(v,pd.Timestamp): v=v.to_pydatetime()
                elif isinstance(v,np.generic): v=v.item()
                vals.append(v)
            ws.append(vals)
    if not tables: wb.create_sheet('Нет данных')
    buffer=io.BytesIO(); wb.save(buffer); return buffer.getvalue()

def decode_arrays(value):
    """Normalize Plotly typed arrays before JSON→Figure and static rendering."""
    if isinstance(value,dict):
        if 'bdata' in value and 'dtype' in value:
            dtype=np.dtype(value['dtype'])
            if dtype.kind not in 'fiu' or dtype.itemsize not in (1,2,4,8):raise ValueError('Неподдерживаемый тип массива графика')
            data=np.frombuffer(base64.b64decode(value['bdata'],validate=True),dtype=dtype)
            if value.get('shape'):
                shape=tuple(int(v.strip()) for v in str(value['shape']).split(','))
                data=data.reshape(shape)
            return data.tolist()
        return {k:decode_arrays(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)):return [decode_arrays(v) for v in value]
    return value


XL_PALETTE=['#4472C4','#ED7D31','#A5A5A5','#FFC000','#5B9BD5','#70AD47','#264478','#9E480E','#636363','#997300','#255E91','#43682B']
XL_TEXT='#595959';XL_GRID='#D9D9D9'

def excel_finish(f,ax,secondary,title,k,date_axis,categorical):
    """Вид стандартной диаграммы Excel (Office 2013+): палитра по порядку рядов, линии 2,25 пт, серый текст,
    только горизонтальные линии сетки (у точечной — и вертикальные), без рамки осей, рамка диаграммы #D9D9D9."""
    from matplotlib.colors import to_hex
    axes=[a for a in (ax,secondary) if a is not None];mapping={}
    def paint(color):
        try:key=to_hex(color,keep_alpha=False).lower()
        except ValueError:return color
        if key in ('#ffffff','#000000'):return color
        return mapping.setdefault(key,XL_PALETTE[len(mapping)%len(XL_PALETTE)])
    for a in axes:
        for line in a.get_lines():
            if len(line.get_xdata())<2 and line.get_marker() in (None,'None',''):continue
            if to_hex(line.get_color()).lower() in ('#000000','#ffffff'):continue
            own=getattr(line,'_own',())       # то, что задано в «Оформлении графиков», вид Excel не меняет
            if 'color' not in own:
                color=paint(line.get_color());hollow=line.get_markerfacecolor()=='white'
                line.set_color(color);line.set_markeredgecolor(color);line.set_markerfacecolor('white' if hollow else color)
            if line.get_linestyle() not in ('None','',' '):
                if 'width' not in own:line.set_linewidth(2.25)
                line.set_solid_capstyle('round');line.set_solid_joinstyle('round')
            if line.get_marker() not in (None,'None',''):
                if 'size' not in own:line.set_markersize(5.5)
                line.set_markeredgewidth(.75)
        for patch in a.patches:patch.set_facecolor(paint(patch.get_facecolor()));patch.set_edgecolor('none')
    for a in axes:
        a.tick_params(axis='both',length=0,labelsize=9*k,labelcolor=XL_TEXT,pad=5);a.xaxis.label.set_size(10*k);a.yaxis.label.set_size(10*k)
        a.xaxis.label.set_color(XL_TEXT);a.yaxis.label.set_color(XL_TEXT)
        for side in ('top','right','left'):a.spines[side].set_visible(False)
        a.spines['bottom'].set_visible(a is ax);a.spines['bottom'].set_color(XL_GRID);a.spines['bottom'].set_linewidth(.75)
    ax.grid(False);ax.grid(True,axis='y',color=XL_GRID,linewidth=.75,linestyle='-')
    if not date_axis and not categorical:ax.grid(True,axis='x',color=XL_GRID,linewidth=.75,linestyle='-')
    if secondary is not None:secondary.grid(False)
    ax.set_title('',loc='left');ax.set_title('',loc='right')
    ax.set_title(title,loc='center',fontsize=14*k,fontweight='normal',color=XL_TEXT,pad=10)
    f.patch.set_facecolor('white');f.patch.set_edgecolor(XL_GRID);f.patch.set_linewidth(1.0)

def x_labels(ax,angle,width_mm,font_pt,date_axis,auto,categorical):
    """Наклон подписей оси X; число подписей подбирается по тому, сколько их помещается в ширину графика при этом наклоне:
    горизонтальная подпись занимает свою ширину, наклонная — проекцию на ось (ширина·cos + высота·sin), вертикальная — только высоту строки.
    Зазор между подписями — полтора размера шрифта. Заданные вручную деления оси X (``auto=False``) не трогаются."""
    import math
    from matplotlib.ticker import MaxNLocator,StrMethodFormatter
    from matplotlib.dates import DateFormatter
    if auto and not categorical:
        theta=math.radians(abs(angle))
        room=width_mm*.8*72/25.4                                  # ширина поля графика без полей осей, пт
        def fits(chars,ticks):                                    # помещаются ли ticks подписей из chars знаков при этом наклоне
            footprint=font_pt*.6*chars*abs(math.cos(theta))+font_pt*1.25*abs(math.sin(theta))+font_pt*1.5
            return ticks*footprint<=room
        if date_axis:
            from matplotlib import dates as md
            lo,hi=ax.get_xlim();days=abs(hi-lo)
            options=[('day',n,n) for n in (1,2,5,10,15)]+[('month',n,30.44*n) for n in (1,2,3,6)]+[('year',n,365.25*n) for n in (1,2,4,5,10,20,50,100)]
            chars={'day':10,'month':7,'year':4}
            unit,n,_=next((o for o in options if fits(chars[o[0]],int(days/o[2])+1)),options[-1])      # самый частый круглый шаг, при котором подписи помещаются
            ax.xaxis.set_major_locator({'day':lambda:md.DayLocator(interval=n),'month':lambda:md.MonthLocator(interval=n),'year':lambda:md.YearLocator(base=n)}[unit]())
            ax.xaxis.set_major_formatter(DateFormatter({'day':'%d.%m.%Y','month':'%m.%Y','year':'%Y'}[unit]))
        else:
            count=next((c for c in range(60,2,-1) if fits(6,c)),3)
            ax.xaxis.set_major_locator(MaxNLocator(nbins=count,steps=[1,2,2.5,5,10]));ax.xaxis.set_major_formatter(StrMethodFormatter('{x:.10g}'))
    ax.tick_params(axis='x',labelrotation=angle)
    for lab in ax.get_xticklabels():lab.set_ha('right' if 0<angle<90 else 'center')
    if angle==90:ax.tick_params(axis='x',labelrotation=90)


def manual_axes(ax,secondary,axes,date_axis,categorical):
    """Ручные границы и деления осей из «Оформления графиков» (``atlas/chart_format.parse_axes``). Сетка основных и дополнительных
    делений начинается с нижней границы оси; не заданное остаётся автоматическим."""
    from matplotlib.ticker import AutoLocator,FixedLocator
    from matplotlib import dates as md
    def ticks(lo,hi,step):
        count=int((hi-lo)/step+1e-9)
        return [lo+i*step for i in range(min(count,600)+1)] if step>0 and count>=0 else []
    def locator(unit,n):
        return {'year':md.YearLocator,'month':md.MonthLocator,'day':md.DayLocator}[unit](**({'base':n} if unit=='year' else {'interval':n}))
    def numeric(spec,get,put,axis):
        lo,hi=get();flip=lo>hi;a,b=sorted((lo,hi))
        a,b=spec.get('min',a),spec.get('max',b)
        if not a<b:return
        put((b,a) if flip else (a,b))
        if 'major' in spec:axis.set_major_locator(FixedLocator(ticks(a,b,spec['major'])))
        elif 'min' in spec or 'max' in spec:axis.set_major_locator(AutoLocator())
        if 'minor' in spec:axis.set_minor_locator(FixedLocator(ticks(a,b,spec['minor'])))
    for name,target in (('y',ax),('y2',secondary)):
        spec=axes.get(name)
        if spec and target is not None:numeric(spec,target.get_ylim,target.set_ylim,target.yaxis)
    spec=axes.get('x')
    if spec and not categorical:
        if date_axis and spec.get('dates'):
            lo,hi=ax.get_xlim()
            a=md.date2num(pd.Timestamp(spec['min']).to_pydatetime()) if 'min' in spec else lo
            b=md.date2num(pd.Timestamp(spec['max']).to_pydatetime()) if 'max' in spec else hi
            if a<b:ax.set_xlim(a,b)
            if 'major' in spec:
                ax.xaxis.set_major_locator(locator(spec['major']['unit'],spec['major']['n']))
                ax.xaxis.set_major_formatter(md.DateFormatter({'year':'%Y','month':'%m.%Y','day':'%d.%m.%Y'}[spec['major']['unit']]))
            if 'minor' in spec:ax.xaxis.set_minor_locator(locator(spec['minor']['unit'],spec['minor']['n']))
        elif not date_axis and not spec.get('dates'):numeric(spec,ax.get_xlim,ax.set_xlim,ax.xaxis)
    if any('minor' in axes.get(k,{}) for k in ('x','y')):
        ax.grid(True,which='minor',color='#f0f2f5',linewidth=.5);ax.tick_params(which='minor',length=2)
    if secondary is not None and 'minor' in axes.get('y2',{}):secondary.tick_params(which='minor',length=2)


def figure_bytes(fig,fmt='png',dpi=300,width_mm=220,height_mm=None,compact=False,font='default',font_size=None,look='default'):
    """``font`` / ``font_size`` — шрифт (``atlas.engine.core.fonts.FONTS``) и основной размер в пт; заголовок и легенда масштабируются вместе с ним.
    ``look='excel'`` — вид диаграммы Excel (Office: Calibri, палитра, тонкие серые горизонтальные линии сетки, легенда внизу); шрифт по умолчанию — Calibri.
    ``compact`` (листы приложений, ``height_mm`` задаёт высоту): график заполняет рисунок, легенда в одну-две строки под осями,
    не больше 6 делений на оси — числа не налезают друг на друга."""
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt
    from matplotlib.ticker import MultipleLocator, StrMethodFormatter
    from matplotlib.dates import AutoDateLocator, DateFormatter
    if fmt not in ('png','svg','pdf'):raise ValueError('Неподдерживаемый формат')
    if dpi not in (150,200,250,300,600,1200):raise ValueError('DPI должен быть 150, 200, 250, 300, 600 или 1200')
    if not 80<=width_mm<=300:raise ValueError('Ширина должна быть от 80 до 300 мм')
    from .fonts import family_for,check
    if look not in ('default','excel'):raise ValueError('Вид графика: default или excel')
    excel=look=='excel'
    if excel and font=='default':font='calibri'
    font,font_size=check(font,font_size)
    base=font_size or (7 if compact else 9);k=base/(7 if compact else 9)
    traces=[t for t in fig.data if t.visible not in (False,'legendonly')]
    legend=[t for t in traces if t.showlegend is not False and t.name]
    if fig.layout.showlegend is False:legend=[]
    longest=max((len(t.name or '') for t in legend),default=0)
    is_gdi=(fig.layout.meta or {}).get('module')=='gdi' or any(isinstance(t.meta,dict) and t.meta.get('module')=='gdi' for t in traces)
    # легенда занимает всю ширину рисунка: число колонок — сколько самых длинных подписей помещается в строку
    lfont=(6.5 if compact else 7)*k
    char_mm=lfont*.5*25.4/72
    avail_mm=width_mm*.94
    wrap_at=max(18,int(avail_mm/char_mm)-8)
    names=[(t.name or '') for t in legend]
    widest=max((min(len(n),wrap_at) for n in names),default=0)
    entry_mm=widest*char_mm+(1.6+.8+1.2)*lfont*25.4/72
    ncols=max(1,min(len(legend) or 1,int(avail_mm/entry_mm)))
    legend_rows=int(np.ceil(len(legend)/ncols))
    if height_mm:pass
    elif is_gdi:height_mm=width_mm*.72+max(0,legend_rows-1)*5*k
    else:height_mm=width_mm*.68+max(0,legend_rows-2)*(9 if longest>=wrap_at else 5)*k
    if fmt=='png' and width_mm*height_mm*(dpi/25.4)**2>90_000_000:raise ValueError('Слишком большой PNG. Уменьшите ширину или DPI.')
    with LOCK,plt.rc_context({'font.family':family_for(font),'font.size':base,'svg.fonttype':'none','pdf.fonttype':42}):
        f,ax=plt.subplots(figsize=(width_mm/25.4,height_mm/25.4),layout='constrained')
        try:
            secondary=ax.twinx() if any(t.yaxis=='y2' for t in traces) else None
            bars=[t for t in traces if t.type=='bar'];boxes=[t for t in traces if t.type=='box'];categories=[];date_axis=fig.layout.xaxis.type=='date'
            numeric_bars=bool(bars) and fig.layout.xaxis.type=='linear'
            if (bars or boxes) and not date_axis and not numeric_bars:
                categories=list(fig.layout.xaxis.categoryarray or [])
                if not categories:categories=list(dict.fromkeys(str(x) for t in (bars or boxes) for x in (t.x if t.x is not None else [t.name])))
            for t in traces:
                target=secondary if secondary is not None and t.yaxis=='y2' else ax
                x=list(decode_arrays(t.x)) if t.x is not None else [];y=np.array(decode_arrays(t.y),dtype=float) if t.y is not None else np.array([])
                label=t.name if t.showlegend is not False else '_nolegend_'
                if t.type=='box':
                    bi=next(i for i,b in enumerate(boxes) if b is t)
                    category=str(t.x[0]) if t.x is not None and len(t.x) else str(t.name)
                    pos=categories.index(category);peers=[b for b in boxes if str(b.x[0] if b.x is not None and len(b.x) else b.name)==category]
                    pi=next(i for i,b in enumerate(peers) if b is t);width=.7/max(1,len(peers));position=pos-.35+width/2+pi*width
                    result=target.boxplot([y[np.isfinite(y)]],positions=[position],widths=width*.9,patch_artist=True,manage_ticks=False,showfliers=t.boxpoints is not False,whis=1.5)
                    color=t.marker.color or '#2563eb'
                    for patch in result['boxes']:patch.set_facecolor(color);patch.set_alpha(.7)
                    if label!='_nolegend_':result['boxes'][0].set_label(label)
                elif t.type=='bar':
                    bi=next(i for i,b in enumerate(bars) if b is t);w=.8/max(1,len(bars))
                    if date_axis:
                        from matplotlib.dates import date2num
                        pos=date2num(pd.to_datetime(x).to_pydatetime())
                    elif numeric_bars:
                        pos=np.asarray(x,dtype=float);raw_width=decode_arrays(t.width)
                        widths=np.asarray(raw_width,dtype=float) if raw_width is not None else np.full(len(pos),.8)
                        w=widths/max(1,len(bars))
                    else:pos=np.array([categories.index(str(v)) for v in x],dtype=float)
                    target.bar(pos-(widths/2 if numeric_bars else .4)+w/2+bi*w,y,width=w,label=label,color=decode_arrays(t.marker.color))
                else:
                    # Plotly JSON serializes dates to strings. Preserve calendar spacing.
                    if x and (date_axis or isinstance(x[0],(pd.Timestamp,np.datetime64,datetime.datetime,datetime.date)) or isinstance(x[0],str) and re.match(r'^\d{4}-\d{2}-\d{2}',x[0])):
                        x=pd.to_datetime(x,errors='raise').to_pydatetime();date_axis=True
                    dash={'solid':'-','dash':'--','dot':':','dashdot':'-.','longdash':'--'}.get(t.line.dash,'-')
                    mode=t.mode or 'lines';color=t.line.color or t.marker.color or '#2563eb'
                    symbol=t.marker.symbol or 'circle';hollow='open' in str(symbol)
                    marker={'square':'s','diamond':'D','triangle-up':'^'}.get(str(symbol).replace('-open',''),'o') if 'markers' in mode else None
                    edge=t.marker.line.color or color
                    face='white' if hollow else t.marker.color or color
                    own=set(t.meta.get('own',())) if isinstance(t.meta,dict) else set()      # задано в «Оформлении графиков» (atlas/chart_format.py)
                    msize=t.marker.size if isinstance(t.marker.size,(int,float)) else None
                    twidth=t.line.width if isinstance(t.line.width,(int,float)) and t.line.width>0 else None
                    ln,=target.plot(x,y,label=label,color=color,linestyle=dash if 'lines' in mode else 'None',marker=marker,
                        markerfacecolor=face,markeredgecolor=edge,markeredgewidth=t.marker.line.width or .7,
                        markersize=msize/2 if msize and 'size' in own else 3,linewidth=twidth if twidth and 'width' in own else 1.2)
                    ln._own=own
            if (bars or boxes) and not date_axis and not numeric_bars:ax.set_xticks(range(len(categories)),categories,rotation=20 if len(categories)>8 else 0)
            title=re.sub('<[^>]+>','',fig.layout.title.text or '')
            fo=(fig.layout.meta or {}).get('format',{}) if isinstance(fig.layout.meta,dict) else {}
            if fo.get('title')=='hide':title=''
            if compact:ax.set_title(title,loc='center',fontsize=8.5*k,fontweight='bold',pad=5)
            else:ax.set_title(title,loc='left',fontsize=12*k,pad=14)
            ax.set_xlabel(fig.layout.xaxis.title.text or '');ax.set_ylabel(fig.layout.yaxis.title.text or '')
            if fig.layout.xaxis.autorange=='reversed' and not date_axis:ax.invert_xaxis()
            if fig.layout.yaxis.autorange=='reversed':ax.invert_yaxis()
            if secondary is not None:
                secondary.set_ylabel(fig.layout.yaxis2.title.text or '')
                if fig.layout.yaxis2.autorange=='reversed':secondary.invert_yaxis()
                secondary.grid(False);secondary.spines['top'].set_visible(False)
            if not date_axis and isinstance(fig.layout.xaxis.dtick,(int,float)) and fig.layout.xaxis.dtick>0:
                lo,hi=ax.get_xlim();step=fig.layout.xaxis.dtick
                if (hi-lo)/step<=150:ax.xaxis.set_major_locator(MultipleLocator(step))
                ax.xaxis.set_major_formatter(StrMethodFormatter('{x:,.0f}' if step>=1 else '{x:.2f}'))
            if isinstance(fig.layout.yaxis.dtick,(int,float)) and fig.layout.yaxis.dtick>0:
                lo,hi=ax.get_ylim();step=fig.layout.yaxis.dtick
                if (hi-lo)/step<=150:ax.yaxis.set_major_locator(MultipleLocator(step))
            if fig.layout.yaxis.rangemode=='tozero':ax.set_ylim(bottom=0)
            if fig.layout.xaxis.rangemode=='tozero' and not date_axis:ax.set_xlim(left=0)
            for axis,layout in [(ax.xaxis,fig.layout.xaxis),(ax.yaxis,fig.layout.yaxis)]:
                axis.grid(layout.showgrid is not False,color='#e4e9ef',linewidth=.7)
            if fig.layout.xaxis.range is not None:ax.set_xlim(*fig.layout.xaxis.range)
            if fig.layout.yaxis.range is not None:ax.set_ylim(*fig.layout.yaxis.range)
            # две оси Y выровнены в Plotly-описании: одинаковое число делений строго друг напротив друга
            for target,layout in ((ax,fig.layout.yaxis),(secondary,fig.layout.yaxis2 if secondary is not None else None)):
                if target is not None and layout.range is not None and layout.dtick and layout.tick0 is not None:
                    a,b=layout.range;target.set_ylim(a,b)
                    from matplotlib.ticker import FixedLocator
                    target.yaxis.set_major_locator(FixedLocator([layout.tick0+i*layout.dtick for i in range(int(round(abs(b-a)/layout.dtick))+1)]))
            if secondary is not None and fig.layout.yaxis2.range is not None:secondary.set_ylim(*fig.layout.yaxis2.range)
            ax.set_axisbelow(True);ax.spines[['top','right']].set_visible(False)
            if excel:excel_finish(f,ax,secondary,title,k,date_axis,bool(bars or boxes) and not numeric_bars)     # до легенды: её значки копируют цвета рядов
            if fo.get('legend')=='hide':legend=[]
            if legend:
                handles,labels=ax.get_legend_handles_labels()
                if secondary is not None:
                    h,l=secondary.get_legend_handles_labels();handles+=h;labels+=l
                if not compact:labels=['\n'.join(textwrap.wrap(str(v),width=wrap_at)) for v in labels]
                if handles:
                    where={'top':'outside upper center','right':'outside center right'}.get(fo.get('legend'),'outside lower center')
                    f.legend(handles,labels,loc=where,ncol=1 if where.endswith('right') else ncols,fontsize=(9 if excel else 6.5 if compact else 7)*k,frameon=False,columnspacing=1.2,handlelength=1.6,labelcolor=XL_TEXT if excel else None)
            if compact:
                from matplotlib.ticker import MaxNLocator
                for axis in (ax.xaxis,ax.yaxis)+((secondary.yaxis,) if secondary is not None else ()):
                    if axis is ax.xaxis and date_axis:continue
                    axis.set_major_locator(MaxNLocator(nbins=6 if width_mm>=120 else 5,steps=[1,2,2.5,5,10]));axis.set_major_formatter(StrMethodFormatter('{x:.10g}'))
                if not excel:ax.grid(True,color='#c9ced6',linewidth=.5,linestyle='--');ax.tick_params(length=2.5,pad=2)
            if date_axis:
                ax.xaxis.set_major_locator(AutoDateLocator(minticks=3,maxticks=5 if compact else 7));ax.xaxis.set_major_formatter(DateFormatter('%m.%Y'))
            if fo.get('axes'):manual_axes(ax,secondary,fo['axes'],date_axis,bool(bars or boxes) and not numeric_bars)
            if fo.get('angle') not in (None,'auto'):x_labels(ax,float(fo['angle']),width_mm,base,date_axis,not (fo.get('axes') or {}).get('x',{}).get('major'),bool(bars or boxes) and not numeric_bars)
            buf=io.BytesIO();f.savefig(buf,format=fmt,dpi=dpi)
            return buf.getvalue()
        finally:plt.close(f)


def export_zip(charts,tables,formats=('svg','pdf'),dpi=300,width_mm=220,metadata=None,progress=None):
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as z:
        for i,(name,fig) in enumerate(charts.items()):
            for fmt in formats: z.writestr(f'charts/{i+1:03}_{safe_name(name)}.{fmt}',figure_bytes(fig,fmt,dpi,width_mm))
            if progress: progress((i+1)/max(1,len(charts)))
        for i,(name,df) in enumerate(tables.items()): z.writestr(f'tables/{i+1:03}_{safe_name(name)}.csv',csv_bytes(df))
        if tables:
            if any(len(d)>1_048_575 for d in tables.values()):
                z.writestr('XLSX_LIMIT.txt','Таблица больше лимита строк Excel. Полные данные сохранены в CSV.')
            else: z.writestr('results.xlsx',xlsx_bytes(tables))
        z.writestr('parameters.json',json.dumps(metadata or {},ensure_ascii=False,indent=2,default=str))
    return buffer.getvalue()
