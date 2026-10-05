"""Правка проблемных строк при импорте: окно строк листа и запись исправлений в исходный файл.

Номер строки — тот же, что в журнале проверки (``Строка``): номер строки листа Excel или строки текстового файла.
Правится только то, что изменил пользователь; остальные ячейки файла не трогаются. XLSX/XLSM — через openpyxl,
текстовые файлы (CSV, TSV, TXT) — построчно, с прежними кодировкой, разделителем и концами строк.
XLS и ODS правке не поддаются: сообщение предлагает сохранить книгу как XLSX.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from typing import Any

import pandas as pd

from app.core import tabular

from .contract import ParamError

MAX_EDITS = 500
NUMBER = re.compile(r'^-?\d+([.,]\d+)?$')
DATE = re.compile(r'^(\d{1,2})\.(\d{1,2})\.(\d{4})$')
NOT_EDITABLE = {'xls': 'Формат XLS не допускает записи: сохраните книгу как XLSX и загрузите её заново.',
                'ods': 'Формат ODS не допускает записи: сохраните книгу как XLSX и загрузите её заново.'}


def show(value: Any) -> str:
    """Значение ячейки так, как его видно в таблице."""
    if value is None or (isinstance(value, float) and value != value) or value is pd.NaT:
        return ''
    if isinstance(value, (dt.datetime, pd.Timestamp)):
        return value.strftime('%d.%m.%Y') if (value.hour, value.minute, value.second) == (0, 0, 0) else value.strftime('%d.%m.%Y %H:%M:%S')
    if isinstance(value, dt.date):
        return value.strftime('%d.%m.%Y')
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return str(value)


def coerce(text: Any) -> Any:
    """Введённый текст → значение ячейки XLSX: число, дата или текст; пусто — пустая ячейка."""
    text = '' if text is None else str(text).strip()
    if not text:
        return None
    if NUMBER.match(text):
        number = float(text.replace(',', '.'))
        return int(number) if number.is_integer() and '.' not in text and ',' not in text else number
    m = DATE.match(text)
    if m:
        try:
            return dt.datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return text
    return text


def window(frame: pd.DataFrame, row: int, radius: int) -> dict[str, Any]:
    total = len(frame)
    if row < 1 or row > total:
        raise ParamError(f'Строки {row} нет на листе: в нём {total} строк.')
    first, last = max(1, row - radius), min(total, row + radius)
    width = frame.shape[1]
    rows = [{'n': n, 'cells': [show(v) for v in frame.iloc[n - 1].tolist()]} for n in range(first, last + 1)]
    return {'row': row, 'first': first, 'last': last, 'total': total, 'width': width, 'rows': rows}


def _encoding(content: bytes, encoding: str) -> str:
    if encoding != 'auto':
        return encoding
    if content.startswith((b'\xff\xfe', b'\xfe\xff')):
        return 'utf-16'
    if content.startswith(b'\xef\xbb\xbf'):
        return 'utf-8-sig'
    try:
        content.decode('utf-8')
        return 'utf-8'
    except UnicodeError:
        return 'cp1251'


def _edit_text(content: bytes, edits: list[tuple[int, int, str]], encoding: str, delimiter: str) -> bytes:
    enc = _encoding(content, encoding)
    text = tabular.decode_text(content, encoding if encoding != 'auto' else enc)
    sep = tabular.resolve_delimiter(text, delimiter)
    if sep == 'whitespace':
        raise ParamError('Файл разделён пробелами: правка недоступна. Укажите разделитель вручную или исправьте файл в редакторе.')
    lines = text.splitlines(keepends=True)
    if len(lines) != sum(1 for _ in csv.reader(io.StringIO(text), delimiter=sep)):
        raise ParamError('В файле есть ячейки с переносами строк: правка недоступна, исправьте файл в редакторе.')
    for row, col, value in edits:
        if row > len(lines):
            raise ParamError(f'Строки {row} нет в файле.')
        line = lines[row - 1]
        end = line[len(line.rstrip('\r\n')):]
        cells = next(csv.reader([line.rstrip('\r\n')], delimiter=sep), [])
        cells += [''] * (col + 1 - len(cells))
        cells[col] = value
        out = io.StringIO()
        csv.writer(out, delimiter=sep, lineterminator='').writerow(cells)
        lines[row - 1] = out.getvalue() + end
    return ''.join(lines).encode(enc)


def _edit_book(content: bytes, name: str, sheet: str, edits: list[tuple[int, int, str]]) -> bytes:
    import openpyxl
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), keep_vba=name.lower().endswith('.xlsm'))
    except Exception as e:
        raise ParamError(f'Книга не открывается для правки: {e}') from None
    if sheet not in book.sheetnames:
        raise ParamError(f'В книге нет листа «{sheet}»')
    ws = book[sheet]
    for row, col, value in edits:
        cell = ws.cell(row=row, column=col + 1)
        try:
            cell.value = coerce(value)
        except AttributeError:
            raise ParamError(f'Ячейка в строке {row} входит в объединённую область: правка недоступна.') from None
        if isinstance(cell.value, dt.datetime):
            cell.number_format = 'DD.MM.YYYY'
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def rewrite(content: bytes, name: str, fmt: str, sheet: str, edits: list, encoding: str, delimiter: str) -> bytes:
    """Файл с применёнными правками. ``edits`` — ``[{row, col, value}]``: row с 1, col с 0."""
    if fmt in NOT_EDITABLE:
        raise ParamError(NOT_EDITABLE[fmt])
    if not edits:
        raise ParamError('Нет изменённых ячеек.')
    if len(edits) > MAX_EDITS:
        raise ParamError(f'За один раз можно изменить не больше {MAX_EDITS} ячеек.')
    try:
        parsed = [(int(e['row']), int(e['col']), '' if e.get('value') is None else str(e['value'])) for e in edits]
    except (KeyError, TypeError, ValueError):
        raise ParamError('Неверные данные правки.') from None
    if any(r < 1 or c < 0 or c > 16383 for r, c, _ in parsed):
        raise ParamError('Неверный адрес ячейки.')
    return _edit_text(content, parsed, encoding, delimiter) if fmt == 'text' else _edit_book(content, name, sheet, parsed)
