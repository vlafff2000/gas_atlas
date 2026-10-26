"""Единый экспорт результата любого модуля.

Графики: ``Chart`` → фигура Plotly → ``atlas.engine.core.export.figure_bytes`` (тот же статический рендер
через Matplotlib, что в 5.8: SVG / PDF / PNG, 300–1200 DPI, без браузера и сети).
Таблицы: ``Table`` → XLSX с подписями столбцов, как на экране.
"""
from __future__ import annotations

import datetime as dt
import re

import math

import numpy as np
import pandas as pd

from atlas.engine.core.config import COLORS
from atlas.engine.core.export import csv_bytes, figure_bytes, safe_name, xlsx_bytes

from .contract import Chart, Series, Table

SYMBOLS = {'circle': 'circle', 'square': 'square', 'diamond': 'diamond', 'triangle': 'triangle-up', 'plus': 'cross', 'x': 'x', 'star': 'star'}
LOOK_DASH = {'solid': 'solid', 'dash': 'dash', 'dot': 'dot'}
FILL_ALPHA = 0.6        # прозрачность заливки области, если не задана вручную
WIDTH_PT = 0.6          # px на экране → пт в выгрузке (2 px по умолчанию = 1,2 пт)
STYLE_AXES = ('x', 'y', 'y2', 'y3', 'y4', 'y5')
FORMATS = {'svg': 'image/svg+xml', 'pdf': 'application/pdf', 'png': 'image/png'}
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def axis_title(label: str, unit: str) -> str:
    return f'{label}, {unit}' if unit else label


# Палитра графиков Атласа 6 (web/src/chartTheme.ts): модули отдают цвета 5.8, выгрузка заменяет их по номеру,
# как экран, — цвета серий в файле совпадают с тем, что видно в окне.
PALETTE = ['#2a78d6', '#eb6834', '#1baf7a', '#eda100', '#e87ba4', '#008300', '#4a3aa7', '#8a5a14']
LEGACY = {old.lower(): new for old, new in zip(COLORS, PALETTE)}


def series_color(color: str) -> str:
    return LEGACY.get(color.lower(), color)


def group_colors(chart: Chart) -> dict[str, str]:
    groups = list(dict.fromkeys(s.group or s.name for s in chart.series))
    return {g: PALETTE[i % len(PALETTE)] for i, g in enumerate(groups)}

TICK_INTERVALS = 5          # одинаковое число основных делений у обеих осей Y (то же число в web/src/ChartView.tsx)
NICE_STEPS = (1, 2, 2.5, 5, 10)


def nice_step(raw: float) -> float:
    """Наименьший «круглый» шаг (1, 2, 2.5, 5 × 10^k), не меньший raw."""
    power = 10.0 ** math.floor(math.log10(raw))
    for m in NICE_STEPS:
        if m * power >= raw * (1 - 1e-9):
            return m * power
    return 10 * power


def aligned_range(lo: float, hi: float, minimum=None, maximum=None, n: int = TICK_INTERVALS):
    """(нижняя граница, шаг) оси ровно из n делений, накрывающей [lo, hi]; заданные границы сохраняются."""
    if minimum is not None and maximum is not None:
        return minimum, (maximum - minimum) / n
    if minimum is not None:
        lo = minimum
    if maximum is not None:
        hi = maximum
    if not hi > lo:
        hi = lo + 1.0
    step = nice_step((hi - lo) / n)
    if minimum is not None:
        return minimum, step
    if maximum is not None:
        return maximum - n * step, step
    while True:
        start = math.floor(lo / step + 1e-9) * step
        if start + n * step >= hi - step * 1e-9:
            return start, step
        step = nice_step(step * 1.0001)


def axis_extent(chart: Chart, key: str):
    lo, hi = math.inf, -math.inf
    stacks: dict[str, list[float]] = {}      # стопка складывается по номеру точки: шкала по сумме, а не по одной серии
    for s in chart.series:
        if (s.axis or 'y') != key or s.kind == 'box':
            continue
        for i, v in enumerate(s.y):
            if isinstance(v, (int, float)) and math.isfinite(v):
                lo, hi = min(lo, v), max(hi, v)
                if s.stack and s.kind != 'points':
                    total = stacks.setdefault(s.stack, [])
                    total.extend([0.0] * (i + 1 - len(total)))
                    total[i] += v
    for total in stacks.values():
        if total:
            lo, hi = min(lo, min(total)), max(hi, max(total))
    return (lo, hi) if lo <= hi else None


def has_bars(chart: Chart) -> bool:
    """На левой оси есть столбцы (в выгрузке столбцы всегда на ней): шкала включает 0."""
    return any(s.kind == 'bar' for s in chart.series)


def align_y_axes(fig, chart: Chart) -> None:
    """Две оси Y: одинаковое число делений, деления строго друг напротив друга — одна основная сетка на обе."""
    specs = {'y': chart.y, **dict(chart.right_axes())}
    if chart.y2 is None or any(a.scale != 'value' for a in specs.values()) or (chart.id.startswith('builder') and any(a.step for a in specs.values())):
        return      # в конструкторе заданный вручную шаг делений важнее общей сетки
    extents = {key: axis_extent(chart, key) for key in specs}
    if all(e is None for e in extents.values()):
        return
    for key, spec in specs.items():
        lo, hi = extents[key] or (0.0, 1.0)      # у оси без данных всё равно те же 5 делений
        if spec.from_zero or (key == 'y' and has_bars(chart)):
            lo, hi = min(lo, 0.0), max(hi, 0.0)
        start, step = aligned_range(lo, hi, spec.minimum, spec.maximum)
        end = start + TICK_INTERVALS * step
        axis = getattr(fig.layout, 'yaxis' if key == 'y' else 'yaxis' + key[1:])
        axis.range = [end, start] if spec.inverse else [start, end]
        axis.autorange = False
        axis.rangemode = 'normal'
        axis.tick0, axis.dtick = start, step


def plot_right_edge(count: int) -> float:
    """Правый край области графика, если справа есть ``count`` шкал: остальное — под подписи шкал."""
    return 0.88 if count == 2 else 1 - 0.06 * count


def right_position(count: int, k: int) -> float:
    """Положение k-й правой шкалы (от 1) по доле ширины: шаг 0,09 при двух шкалах, 0,06 при больше."""
    if count == 2:
        return 0.88 + 0.09 * (k - 1)
    return plot_right_edge(count) + 0.06 * (k - 1)


def rgba(color: str, alpha: float) -> str:
    c = color.lstrip('#')
    if len(c) != 6:
        return color
    return f'rgba({int(c[0:2], 16)},{int(c[2:4], 16)},{int(c[4:6], 16)},{alpha})'


def look_meta(look: dict) -> dict:
    """Что задано вручную (``Series.look``): статический рендер берёт эти значения вместо своих по умолчанию (``own``)."""
    own = [name for name, key in (('width', 'width'), ('size', 'size'), ('alpha', 'opacity')) if key in look]
    return {'own': own}


def style_meta(chart: Chart) -> dict:
    """Ручное оформление графика и осей для статического рендера (``atlas/engine/core/export_style.py``)."""
    axes = {}
    for name in STYLE_AXES:
        spec = chart.x if name == 'x' else chart.axis_at(name) if name in dict(chart.right_axes()) or name == 'y' else None
        if spec is None:
            continue
        item = {k: v for k, v in spec.style.items() if v not in (None, '')}
        if spec.minimum is not None:
            item['min'] = spec.minimum
        if spec.maximum is not None:
            item['max'] = spec.maximum
        if spec.step:
            item['step'] = spec.step
        if spec.color:
            item['color'] = spec.color
        item['scale'] = spec.scale
        axes[name] = item
    return {'chart': dict(chart.style), 'axes': axes}


NARROW = 5      # сегмент стопки уже 5 % высоты: число не помещается внутрь (как на экране, ChartView.tsx)


def segment_marks(fig, chart: Chart, palette: dict) -> None:
    """Узкие сегменты стопки с подписями: выноска в рамке цвета ряда (по очереди вправо и влево) или точка; таблица под графиком."""
    import plotly.graph_objects as go
    bars = [s for s in chart.series if s.kind == 'bar' and s.segments is not None]
    if not bars:
        return
    below = np.zeros(len(bars[0].x))
    used = [0] * len(bars[0].x)
    for s in bars:
        color = (s.look or {}).get('color') or (series_color(s.color) if s.color else palette[s.group or s.name])
        for j, (x, y, text) in enumerate(zip(s.x, s.y, s.segments)):
            y = 0.0 if y is None else float(y)
            middle = below[j] + y / 2
            below[j] += y
            if not text or y >= NARROW:
                continue
            if chart.callouts == 'dot':
                fig.add_annotation(x=x, y=middle, text='●', showarrow=False, font={'size': 9, 'color': '#111'})
                continue
            side = 1 if used[j] % 2 == 0 else -1
            used[j] += 1
            fig.add_annotation(x=x, y=middle, text=text, showarrow=True, ax=side * 44, ay=0, axref='pixel', ayref='pixel', arrowhead=0,
                               arrowwidth=0.8, arrowcolor='#888', bordercolor=color, borderwidth=1.5, borderpad=2, bgcolor='white',
                               font={'size': 12, 'color': '#111'})
    if chart.table is not None:
        t = chart.table
        colors = [(s.look or {}).get('color') or (series_color(s.color) if s.color else palette[s.group or s.name]) for s in bars] if t.swatches else []
        colors = colors + ['white'] * (len(t.rows) - len(colors))
        fig.add_trace(go.Table(domain={'x': [0.0, 1.0], 'y': [0.0, 0.36]}, columnwidth=[2.4] + [1] * len(t.columns),
                               header={'values': [''] + list(t.columns), 'fill_color': 'white', 'align': 'center', 'font': {'size': 12}},
                               cells={'values': [list(t.rows)] + [[row[j] for row in t.cells] for j in range(len(t.columns))],
                                      'fill_color': [colors] + [['white'] * len(t.rows)] * len(t.columns), 'align': ['left'] + ['center'] * len(t.columns),
                                      'font': {'size': 12}, 'height': 22}))
        fig.update_layout(yaxis={'domain': [0.42, 1.0]}, height=720)


def to_plotly(chart: Chart):
    import plotly.graph_objects as go
    palette = group_colors(chart)
    fig = go.Figure()
    stacked_bars = any(s.kind == 'bar' and s.stack for s in chart.series)
    combo = chart.id.startswith('builder')      # виды конструктора: область, ступенчатая, стопка (atlas/engine/core/export.py читает meta['combo'])
    scale_title: str | None = None
    for n, s in enumerate(chart.series):
        look = s.look or {}
        color = look.get('color') or (series_color(s.color) if s.color else palette[s.group or s.name])
        if s.kind == 'box':
            # Пять чисел ящика как значения: квартили по ним совпадают, усы — крайние значения в пределах 1,5 IQR.
            first = True
            for category, stats in zip(s.x, s.y):
                if stats is None:
                    continue
                fig.add_trace(go.Box(x=[str(category)] * len(stats), y=list(stats), name=s.name,
                                     showlegend=s.legend and first, marker_color=color, boxpoints=False))
                first = False
        elif s.kind == 'bar':
            extra = {'meta': {'value_labels': True}} if look.get('valueLabels') else {}     # подписи значений: рисует atlas/engine/core/export.py
            if combo:       # конструктор: столбцы на своей оси, рядом или стопкой по имени, дата остаётся датой
                extra = dict(yaxis=s.axis, offsetgroup=s.stack or f'bar{n}', opacity=look.get('opacity'),
                             meta={'combo': True, 'stack': s.stack, **extra.get('meta', {})})
            if s.segments is not None:       # стопка с числами внутри сегментов: узкие — выноской или точкой (``segment_marks``)
                wide = [t if (y is not None and y >= NARROW) else '' for t, y in zip(s.segments, s.y)]
                extra.update(text=wide, textposition='inside', insidetextanchor='middle', textfont={'size': 13},
                             marker={'color': color, 'line': {'color': '#222', 'width': 0.4}})
                extra.pop('marker_color', None)
            fig.add_trace(go.Bar(x=list(s.x) if stacked_bars or combo else [str(v) for v in s.x], y=list(s.y), name=s.name,
                                 showlegend=s.legend, **({} if 'marker' in extra else {'marker_color': color}), **extra))
        elif s.kind == 'line':
            dash = LOOK_DASH.get(look.get('dash', ''), s.dash or ('dash' if s.dashed else 'solid'))
            shown = look.get('line', True)
            markers = look.get('marker', s.markers)
            extra = {'meta': look_meta(look)}
            line_extra: dict = {}
            if combo:
                extra['meta'].update(combo=True, stack=s.stack, fill=s.fill, step=s.step)
                if s.step:
                    line_extra['shape'] = 'hv'      # значение держится до следующей точки, как на экране
                if s.fill:
                    alpha = look.get('opacity', FILL_ALPHA)
                    extra['fillcolor'] = rgba(color, alpha)
                    extra['meta']['alpha'] = alpha
                    extra['fill'] = 'tonexty' if s.stack else 'tozeroy'
            if markers:
                shape = SYMBOLS[look.get('shape', s.symbol)]
                mcolor = look.get('markerColor') or color
                extra['marker'] = {'color': 'white' if s.hollow else mcolor, 'symbol': shape, 'size': look.get('size', 6),
                                   'line': {'color': mcolor, 'width': 1.8 if s.hollow or shape in ('cross', 'x') else 0}}
            if s.stack:      # линии одной стопки складываются по значению X (дате), как на экране
                extra['stackgroup'] = s.stack
            if not shown and not markers and not s.fill:
                extra['visible'] = False         # у ряда нет ни линии, ни маркеров — в файл не попадает
            if not shown and s.fill:
                line_extra['width'] = 0           # область без линии: остаётся заливка
            mode = ('lines' if shown or s.fill else '') + ('+' if (shown or s.fill) and markers else '') + ('markers' if markers else '')
            fig.add_trace(go.Scatter(x=list(s.x), y=list(s.y), mode=mode or 'lines',
                                     name=s.name, showlegend=s.legend, yaxis=s.axis,
                                     connectgaps=False, opacity=s.opacity,
                                     line={'color': color, 'dash': dash, 'width': round(look['width'] * WIDTH_PT, 2) if 'width' in look else s.width or 2,
                                           **line_extra},
                                     **extra))
        else:
            shape = SYMBOLS[look.get('shape', s.symbol)]
            symbol = shape + ('-open' if s.hollow else '')
            marker = {'color': look.get('markerColor') or color, 'symbol': symbol, 'size': look.get('size', 9),
                      'opacity': s.opacity,
                      'line': {'color': look.get('markerColor') or color, 'width': 2 if s.hollow or shape in ('cross', 'x') else 0}}
            if s.values is not None:         # точки окрашены третьим значением: общая шкала цвета на весь график
                marker.update(color=[float('nan') if v is None or v != v else float(v) for v in s.values], coloraxis='coloraxis')
                marker['line'] = {'width': 0}
                scale_title = scale_title if scale_title is not None else s.values_label
            fig.add_trace(go.Scatter(x=list(s.x), y=list(s.y), mode='markers', name=s.name, showlegend=s.legend,
                                     yaxis=s.axis, meta=look_meta(look), opacity=s.opacity, marker=marker))
    segment_marks(fig, chart, palette)
    module = chart.id.split('-', 1)[0]
    meta: dict = {'module': module}
    if chart.id.startswith('builder') or chart.style or any(a.style for a in (chart.x, chart.y, *(a for _, a in chart.right_axes()))):
        meta['style'] = style_meta(chart)
        if chart.style.get('legend'):
            meta['format'] = {'legend': chart.style['legend']}
    fig.update_layout(title={'text': chart.title}, meta=meta, barmode='stack' if stacked_bars else 'group')
    if scale_title is not None:
        fig.update_layout(coloraxis={'colorscale': 'Viridis', 'colorbar': {'title': {'text': scale_title}}})
    axes = [(fig.layout.xaxis, chart.x), (fig.layout.yaxis, chart.y)]
    right = chart.right_axes()
    if len(right) > 1:      # место под подписи нескольких правых шкал; одна правая шкала — как раньше, на всю ширину
        fig.update_layout(xaxis={'domain': [0, plot_right_edge(len(right))]})
    for k, (name, spec) in enumerate(right, start=1):
        # первая правая шкала у графика, следующие — снаружи от предыдущей (как на экране)
        layout = {'overlaying': 'y', 'side': 'right', 'showgrid': False}
        if len(right) > 1:
            layout.update(anchor='free', position=right_position(len(right), k))
        fig.update_layout(**{'yaxis' + name[1:]: layout})
        axes.append((getattr(fig.layout, 'yaxis' + name[1:]), spec))
    for n, (axis, spec) in enumerate(axes):
        axis.title = {'text': axis_title(spec.label, spec.unit)}
        if spec.color:      # правая шкала под один ряд: подписи, название и линия оси — цветом ряда
            axis.title = {'text': axis_title(spec.label, spec.unit), 'font': {'color': spec.color}}
            axis.tickfont = {'color': spec.color}
            axis.linecolor, axis.showline = spec.color, True
        axis.type = {'time': 'date', 'log': 'log', 'category': 'category'}.get(spec.scale, 'linear')
        if spec.scale == 'category' and spec.categories:
            axis.categoryorder = 'array'
            axis.categoryarray = list(spec.categories)
        if (spec.from_zero or (n == 1 and has_bars(chart))) and spec.scale != 'category':
            axis.rangemode = 'tozero'
        if spec.step and spec.scale in ('value', 'log'):
            axis.dtick = spec.step
        if spec.minimum is not None and spec.maximum is not None and spec.scale == 'value':
            axis.range = [spec.minimum, spec.maximum]
        if spec.inverse:
            axis.autorange = 'reversed'
    align_y_axes(fig, chart)
    return fig


def _cell(value):
    """Значение для ячейки Excel: даты — datetime, пустые и нечисловые — None, numpy — обычные числа."""
    if value is None:
        return None
    if isinstance(value, (pd.Timestamp, np.datetime64)):
        return None if pd.isna(value) else pd.Timestamp(value).to_pydatetime()
    if isinstance(value, (np.floating, float)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, np.integer):
        return int(value)
    return value


def chart_xlsx(chart: Chart) -> tuple[bytes, str, str]:
    """XLSX графика: лист «Данные» (все точки каждого ряда, подсказки) и лист «График» — тот же график средствами Excel.

    Ряды-столбцы по категориям (общая первая колонка), остальные — парами «X, Y»; ряды правой оси — вторым графиком поверх.
    Ряды-«ящики» в Excel не рисуются: их числа остаются на листе «Данные».
    """
    import io
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference, ScatterChart, Series as XSeries
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.styles import Alignment, Font

    xl, yl = axis_title(chart.x.label, chart.x.unit), axis_title(chart.y.label, chart.y.unit)
    right = chart.right_axes()
    unit_of = {'y': chart.y.unit, **{name: spec.unit for name, spec in right}}
    title_of = {'y': yl, **{name: axis_title(spec.label, spec.unit) for name, spec in right}}
    wb = Workbook()
    data = wb.active
    data.title = 'Данные'
    category = chart.x.scale == 'category'
    plotted: list[tuple[Series, int, int, int, int]] = []       # ряд, колонка X, колонка Y, первая и последняя строка
    if category:
        cats = list(chart.x.categories or dict.fromkeys(str(v) for s in chart.series for v in s.x))
        data.cell(1, 1, xl)
        for r, c in enumerate(cats, 2):
            data.cell(r, 1, c)
        for k, s in enumerate(chart.series, 2):
            data.cell(1, k, f'{s.name}, {unit_of.get(s.axis, chart.y.unit)}'.rstrip(', '))
            by = {str(x): y for x, y in zip(s.x, s.y)}
            for r, c in enumerate(cats, 2):
                data.cell(r, k, _cell(by.get(c)))
            plotted.append((s, 1, k, 2, len(cats) + 1))
        width = len(chart.series) + 1
    else:
        col = 1
        for s in chart.series:
            data.cell(1, col, f'{s.name} · {xl}')
            data.cell(1, col + 1, f'{s.name} · {title_of.get(s.axis, yl)}')
            extra = [s.labels] if s.labels is not None else []
            if extra:
                data.cell(1, col + 2, f'{s.name} · подсказка')
            for r, (x, y) in enumerate(zip(s.x, s.y), 2):
                cx, cy = _cell(x), _cell(y)
                data.cell(r, col, cx)
                data.cell(r, col + 1, cy)
                if cx is not None and hasattr(cx, 'year'):
                    data.cell(r, col).number_format = 'DD.MM.YYYY'
                if extra and r - 2 < len(s.labels):
                    data.cell(r, col + 2, s.labels[r - 2])
            plotted.append((s, col, col + 1, 2, len(s.y) + 1))
            col += 3 if extra else 2
        width = col - 1
    for c in range(1, width + 1):
        data.cell(1, c).font = Font(bold=True)
        data.cell(1, c).alignment = Alignment(wrap_text=True, vertical='top')
        data.column_dimensions[data.cell(1, c).column_letter].width = 22
    data.freeze_panes = 'A2'

    def build(items, secondary: bool):
        if category:
            part = BarChart()
            part.type = 'col'
            if any(s.stack for s, *_ in items):
                part.grouping, part.overlap = 'stacked', 100
        else:
            part = ScatterChart()
            part.style = 13
        for n, (s, cx, cy, r0, r1) in enumerate(items):
            values = Reference(data, min_col=cy, min_row=1, max_row=r1)
            if category:
                one = XSeries(values, title_from_data=True)
            else:
                one = XSeries(Reference(data, min_col=cy, min_row=r0, max_row=r1), Reference(data, min_col=cx, min_row=r0, max_row=r1),
                              title=s.name)
                if s.kind == 'points':
                    one.marker.symbol = 'circle'
                    one.marker.size = 5
                    one.graphicalProperties.line.noFill = True
                else:
                    one.marker.symbol = 'none'
                    one.smooth = False
                    one.graphicalProperties.line.width = 19050
            color = ((s.look or {}).get('color') or (series_color(s.color) if s.color else PALETTE[(n + (3 if secondary else 0)) % len(PALETTE)])).lstrip('#')
            if color:
                if category:
                    one.graphicalProperties.solidFill = color
                elif s.kind == 'points':
                    one.marker.graphicalProperties = GraphicalProperties(solidFill=color)
                else:
                    one.graphicalProperties.line.solidFill = color
            part.series.append(one)
        return part

    drawable = [item for item in plotted if item[0].kind != 'box']
    if drawable:
        main = build([i for i in drawable if i[0].axis == 'y'] or drawable, False)
        main.title = chart.title
        main.x_axis.title, main.y_axis.title = xl, yl
        main.x_axis.delete = main.y_axis.delete = False
        main.height, main.width = 11, 20
        if not category and chart.x.scale == 'time':
            main.x_axis.number_format = 'DD.MM.YYYY'
        # в Excel у диаграммы одна правая ось: все правые шкалы идут на неё, единицы каждого ряда — в заголовке листа «Данные»
        second = [i for i in drawable if i[0].axis != 'y'] if right and any(i[0].axis == 'y' for i in drawable) else []
        if second:
            secondary = build(second, True)
            secondary.y_axis.axId = 200
            secondary.y_axis.title = ' · '.join(title_of[name] for name, _ in right)
            secondary.y_axis.crosses = 'max'
            secondary.y_axis.delete = False
            main += secondary
        sheet = wb.create_sheet('График')
        sheet.add_chart(main, 'A1')
    return_name = f'{safe_name(chart.title)}.xlsx'
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), return_name, XLSX


def chart_csv(chart: Chart) -> tuple[bytes, str, str]:
    """Данные графика: по строке на точку — ряд, ось, X, Y (CSV как в 5.8: «;», запятая в дробях, UTF-8 с BOM)."""
    frames = []
    for s in chart.series:
        if s.kind == 'box' or len(s.x) != len(s.y):
            continue
        frames.append(pd.DataFrame({'Ряд': s.name, 'Ось': s.axis, 'X': list(s.x), 'Y': list(s.y)}))
    frame = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=['Ряд', 'Ось', 'X', 'Y'])
    for col in ('X', 'Y'):
        if frame[col].map(lambda v: isinstance(v, (pd.Timestamp, dt.datetime, dt.date))).any():
            frame[col] = frame[col].map(lambda v: pd.Timestamp(v).strftime('%Y-%m-%d') if v is not None and not pd.isna(v) else '')
    return csv_bytes(frame), f'{safe_name(chart.title)}.csv', 'text/csv; charset=utf-8'


def chart_file(chart: Chart, fmt: str, dpi: int, font: str | None = None, font_size=None) -> tuple[bytes, str, str]:
    if fmt == 'xlsx':
        return chart_xlsx(chart)
    if fmt == 'csv':
        return chart_csv(chart)
    if fmt not in FORMATS:
        raise ValueError('Формат графика: svg, pdf, png, xlsx или csv')
    return figure_bytes(to_plotly(chart), fmt, dpi, font=font or 'default', font_size=font_size), f'{safe_name(chart.title)}.{fmt}', FORMATS[fmt]


def column_titles(table: Table) -> list[str]:
    """Подписи колонок; у таблицы с объединённой шапкой — с подписями групп (в плоском файле иначе не различить колонки)."""
    titles = []
    for k, c in enumerate(table.columns):
        groups, at = [], 0
        for row in table.header:
            for label, span in row:
                if at <= k < at + span and label:
                    groups.append(label)
                at += span
            at = 0
        titles.append(axis_title(' · '.join(groups + [c.label]), c.unit))
    return titles


def table_frame(table: Table) -> pd.DataFrame:
    keys = [c.key for c in table.columns]
    return table.frame[keys].set_axis(column_titles(table), axis=1)


def sheet_formula(table: Table, key: str, at: int, first: int) -> str | None:
    """Формула Excel для колонки ``key`` в строке ``at`` (первая строка данных — ``first``) или None, если её нет."""
    from openpyxl.utils import get_column_letter
    template = table.formulas.get(key)
    if not template:
        return None
    letter = {c.key: get_column_letter(k) for k, c in enumerate(table.columns, 1)}
    if '@prev' in template and at == first:
        return None

    last = first + len(table.frame) - 1

    def cell(match):
        name, mode = match.group(1), match.group(2)
        if mode == '@col':
            return f'${letter[name]}${first}:${letter[name]}${last}'
        return f'{letter[name]}{at - 1 if mode else at}'
    return '=' + re.sub(r'\{(\w+)(@prev|@col)?\}', cell, template)


def header_sheet(ws, table: Table) -> None:
    """Лист с объединённой шапкой, как в «Сводная_транспон2.xlsx»: группы сверху, подписи колонок внизу."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from atlas.engine.core.export import safe_table
    side = Side(style='thin', color='000000')
    fill, rows = PatternFill('solid', fgColor='D9E1F2'), len(table.header) + 1
    for r, row in enumerate(table.header, 1):
        at = 1
        for label, span in row:
            if label:
                ws.cell(r, at, label)
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=at, end_row=r, end_column=at + span - 1)
            at += span
    for k, c in enumerate(table.columns, 1):
        ws.cell(rows, k, axis_title(c.label, c.unit))
        if not table.header[0][0][0] and k <= table.header[0][0][1]:    # колонка без группы: во всю высоту шапки
            ws.merge_cells(start_row=1, start_column=k, end_row=rows, end_column=k)
    for r in range(1, rows + 1):
        for k in range(1, len(table.columns) + 1):
            cell = ws.cell(r, k)
            cell.font, cell.fill, cell.border = Font(bold=True), fill, Border(left=side, right=side, top=side, bottom=side)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    body = safe_table(table.frame[[c.key for c in table.columns]])
    for r, values in enumerate(body.itertuples(index=False, name=None), rows + 1):
        for k, (v, c) in enumerate(zip(values, table.columns), 1):
            v = None if pd.isna(v) else v.item() if hasattr(v, 'item') else v
            formula = sheet_formula(table, c.key, r, rows + 1)
            cell = ws.cell(r, k, formula if formula else None if c.key in table.formulas else v)
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if c.kind == 'number' and c.decimals is not None:
                cell.number_format = '0.' + '0' * c.decimals if c.decimals else '0'
    ws.freeze_panes = ws.cell(rows + 1, 1)
    ws.column_dimensions['B'].width = 38


def plain_sheet(ws, table: Table) -> None:
    """Таблица без объединённой шапки в том же оформлении: жирная шапка с заливкой, рамки, форматы чисел и дат, ширина по тексту."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    side = Side(style='thin', color='000000')
    border, fill = Border(left=side, right=side, top=side, bottom=side), PatternFill('solid', fgColor='D9E1F2')
    frame = table.frame[[c.key for c in table.columns]]
    titles = column_titles(table)
    for k, title in enumerate(titles, 1):
        cell = ws.cell(1, k, title)
        cell.font, cell.fill, cell.border = Font(bold=True), fill, border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    widths = [min(max(len(t) * 0.55, 10), 40) for t in titles]
    for r, values in enumerate(frame.itertuples(index=False, name=None), 2):
        for k, (v, c) in enumerate(zip(values, table.columns), 1):
            v = None if v is None or (not isinstance(v, (list, tuple)) and pd.isna(v)) else v
            v = v.to_pydatetime() if isinstance(v, pd.Timestamp) else v.item() if hasattr(v, 'item') else v
            formula = sheet_formula(table, c.key, r, 2)
            cell = ws.cell(r, k, formula if formula else None if c.key in table.formulas else v)
            cell.border = border
            if c.kind == 'number' and c.decimals is not None:
                cell.number_format = '0.' + '0' * c.decimals if c.decimals else '0'
            elif c.kind == 'date' or hasattr(v, 'year'):
                cell.number_format = 'DD.MM.YYYY'
            elif isinstance(v, str):
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                widths[k - 1] = min(max(widths[k - 1], len(v) * 0.9), 60)
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = round(w, 1)
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = 'A2'
    if table.note:
        ws.cell(len(frame) + 3, 1, table.note).font = Font(italic=True)


def tables_file(tables: list[Table], name: str) -> tuple[bytes, str, str]:
    """XLSX с листом на таблицу: оформление как на экране (шапка, форматы, ширина), сверху вниз — как в отчёте."""
    import io
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for i, t in enumerate(tables, 1):
        base = re.sub(r'[\[\]:*?/\\]', '_', t.title)[:27]
        title = base if base not in used else f'{base}_{i}'
        used.add(title)
        ws = wb.create_sheet(title)
        if t.header:
            header_sheet(ws, t)
        else:
            plain_sheet(ws, t)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), f'{safe_name(name)}.xlsx', XLSX


def table_csv(table: Table) -> tuple[bytes, str, str]:
    """CSV как в 5.8: разделитель «;», запятая в дробях, UTF-8 с BOM — открывается в Excel."""
    return csv_bytes(table_frame(table)), f'{safe_name(table.title)}.csv', 'text/csv; charset=utf-8'
