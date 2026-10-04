"""Local static exports; no Chrome/CDN required. Plotly traces → Matplotlib."""
import base64
import io
import json
import re
import zipfile
import threading
import textwrap
import datetime
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

LOCK=threading.RLock()

def safe_name(name):
    return re.sub(r'[^\w .-]','_',str(name)).strip(' .')[:100] or 'chart'

def safe_table(df):
    d=df.copy()
    for col in d.select_dtypes(include=['object','string']):
        d[col]=d[col].map(lambda v:"'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v)
    return d

def csv_bytes(df):
    return safe_table(df).to_csv(index=False,sep=';',decimal=',').encode('utf-8-sig')

def xlsx_bytes(tables):
    wb=Workbook(write_only=True)
    from openpyxl.cell import WriteOnlyCell
    for i,(name,df) in enumerate(tables.items()):
        name=re.sub(r'[\[\]:*?/\\]','_',name)[:25]+'_'+str(i+1)
        ws=wb.create_sheet(name); ws.freeze_panes='A2'
        headers=[]
        for col in df:
            c=WriteOnlyCell(ws,str(col)); c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='243247'); headers.append(c)
        ws.append(headers)
        for row in safe_table(df).itertuples(index=False,name=None):
            vals=[]
            for v in row:
                if pd.isna(v): v=None
                elif isinstance(v,pd.Timestamp): v=v.to_pydatetime()
                elif isinstance(v,np.generic): v=v.item()
                vals.append(v)
            ws.append(vals)
    if not tables: wb.create_sheet('Нет данных')
    buffer=io.BytesIO(); wb.save(buffer); return buffer.getvalue()

def decode_arrays(value):
    """Normalize Plotly typed arrays before JSON→Figure and static rendering."""
    if isinstance(value,dict):
        if 'bdata' in value and 'dtype' in value:
            dtype=np.dtype(value['dtype'])
            if dtype.kind not in 'fiu' or dtype.itemsize not in (1,2,4,8):raise ValueError('Неподдерживаемый тип массива графика')
            data=np.frombuffer(base64.b64decode(value['bdata'],validate=True),dtype=dtype)
            if value.get('shape'):
                shape=tuple(int(v.strip()) for v in str(value['shape']).split(','))
                data=data.reshape(shape)
            return data.tolist()
        return {k:decode_arrays(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)):return [decode_arrays(v) for v in value]
    return value


def figure_bytes(fig,fmt='png',dpi=300,width_mm=220):
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt
    from matplotlib.ticker import MultipleLocator, StrMethodFormatter
    from matplotlib.dates import AutoDateLocator, DateFormatter
    if fmt not in ('png','svg','pdf'):raise ValueError('Неподдерживаемый формат')
    if dpi not in (150,300,600,1200):raise ValueError('DPI должен быть 150, 300, 600 или 1200')
    if not 80<=width_mm<=300:raise ValueError('Ширина должна быть от 80 до 300 мм')
    traces=[t for t in fig.data if t.visible not in (False,'legendonly')]
    legend=[t for t in traces if t.showlegend is not False and t.name]
    if fig.layout.showlegend is False:legend=[]
    longest=max((len(t.name or '') for t in legend),default=0)
    is_gdi=(fig.layout.meta or {}).get('module')=='gdi' or any(isinstance(t.meta,dict) and t.meta.get('module')=='gdi' for t in traces)
    ncols=(max(1,int(width_mm/70)) if is_gdi else 1 if width_mm<150 else (3 if longest<22 else 2))
    legend_rows=int(np.ceil(len(legend)/ncols))
    height_mm=width_mm*.72+max(0,legend_rows-1)*5 if is_gdi else width_mm*.68+max(0,legend_rows-2)*(9 if longest>=40 else 5)
    if fmt=='png' and width_mm*height_mm*(dpi/25.4)**2>90_000_000:raise ValueError('Слишком большой PNG. Уменьшите ширину или DPI.')
    with LOCK,plt.rc_context({'font.family':'DejaVu Sans','font.size':9,'svg.fonttype':'none','pdf.fonttype':42}):
        f,ax=plt.subplots(figsize=(width_mm/25.4,height_mm/25.4),layout='constrained')
        try:
            secondary=ax.twinx() if any(t.yaxis=='y2' for t in traces) else None
            bars=[t for t in traces if t.type=='bar'];boxes=[t for t in traces if t.type=='box'];categories=[];date_axis=fig.layout.xaxis.type=='date'
            numeric_bars=bool(bars) and fig.layout.xaxis.type=='linear'
            if (bars or boxes) and not date_axis and not numeric_bars:
                categories=list(fig.layout.xaxis.categoryarray or [])
                if not categories:categories=list(dict.fromkeys(str(x) for t in (bars or boxes) for x in (t.x if t.x is not None else [t.name])))
            for t in traces:
                target=secondary if secondary is not None and t.yaxis=='y2' else ax
                x=list(decode_arrays(t.x)) if t.x is not None else [];y=np.array(decode_arrays(t.y),dtype=float) if t.y is not None else np.array([])
                label=t.name if t.showlegend is not False else '_nolegend_'
                if t.type=='box':
                    bi=next(i for i,b in enumerate(boxes) if b is t)
                    category=str(t.x[0]) if t.x is not None and len(t.x) else str(t.name)
                    pos=categories.index(category);peers=[b for b in boxes if str(b.x[0] if b.x is not None and len(b.x) else b.name)==category]
                    pi=next(i for i,b in enumerate(peers) if b is t);width=.7/max(1,len(peers));position=pos-.35+width/2+pi*width
                    result=target.boxplot([y[np.isfinite(y)]],positions=[position],widths=width*.9,patch_artist=True,manage_ticks=False,showfliers=t.boxpoints is not False,whis=1.5)
                    color=t.marker.color or '#2563eb'
                    for patch in result['boxes']:patch.set_facecolor(color);patch.set_alpha(.7)
                    if label!='_nolegend_':result['boxes'][0].set_label(label)
                elif t.type=='bar':
                    bi=next(i for i,b in enumerate(bars) if b is t);w=.8/max(1,len(bars))
                    if date_axis:
                        from matplotlib.dates import date2num
                        pos=date2num(pd.to_datetime(x).to_pydatetime())
                    elif numeric_bars:
                        pos=np.asarray(x,dtype=float);raw_width=decode_arrays(t.width)
                        widths=np.asarray(raw_width,dtype=float) if raw_width is not None else np.full(len(pos),.8)
                        w=widths/max(1,len(bars))
                    else:pos=np.array([categories.index(str(v)) for v in x],dtype=float)
                    target.bar(pos-(widths/2 if numeric_bars else .4)+w/2+bi*w,y,width=w,label=label,color=decode_arrays(t.marker.color))
                else:
                    # Plotly JSON serializes dates to strings. Preserve calendar spacing.
                    if x and (date_axis or isinstance(x[0],(pd.Timestamp,np.datetime64,datetime.datetime,datetime.date)) or isinstance(x[0],str) and re.match(r'^\d{4}-\d{2}-\d{2}',x[0])):
                        x=pd.to_datetime(x,errors='raise').to_pydatetime();date_axis=True
                    dash={'solid':'-','dash':'--','dot':':','dashdot':'-.','longdash':'--'}.get(t.line.dash,'-')
                    mode=t.mode or 'lines';color=t.line.color or t.marker.color or '#2563eb'
                    symbol=t.marker.symbol or 'circle';hollow='open' in str(symbol)
                    marker={'square':'s','diamond':'D','triangle-up':'^'}.get(str(symbol).replace('-open',''),'o') if 'markers' in mode else None
                    edge=t.marker.line.color or color
                    face='white' if hollow else t.marker.color or color
                    target.plot(x,y,label=label,color=color,linestyle=dash if 'lines' in mode else 'None',marker=marker,
                        markerfacecolor=face,markeredgecolor=edge,markeredgewidth=t.marker.line.width or .7,
                        markersize=3,linewidth=1.2)
            if (bars or boxes) and not date_axis and not numeric_bars:ax.set_xticks(range(len(categories)),categories,rotation=20 if len(categories)>8 else 0)
            ax.set_title(re.sub('<[^>]+>','',fig.layout.title.text or ''),loc='left',fontsize=12,pad=14)
            ax.set_xlabel(fig.layout.xaxis.title.text or '');ax.set_ylabel(fig.layout.yaxis.title.text or '')
            if fig.layout.xaxis.autorange=='reversed' and not date_axis:ax.invert_xaxis()
            if fig.layout.yaxis.autorange=='reversed':ax.invert_yaxis()
            if secondary is not None:
                secondary.set_ylabel(fig.layout.yaxis2.title.text or '')
                if fig.layout.yaxis2.autorange=='reversed':secondary.invert_yaxis()
                secondary.grid(False);secondary.spines['top'].set_visible(False)
            if not date_axis and isinstance(fig.layout.xaxis.dtick,(int,float)) and fig.layout.xaxis.dtick>0:
                lo,hi=ax.get_xlim();step=fig.layout.xaxis.dtick
                if (hi-lo)/step<=150:ax.xaxis.set_major_locator(MultipleLocator(step))
                ax.xaxis.set_major_formatter(StrMethodFormatter('{x:,.0f}' if step>=1 else '{x:.2f}'))
            if isinstance(fig.layout.yaxis.dtick,(int,float)) and fig.layout.yaxis.dtick>0:
                lo,hi=ax.get_ylim();step=fig.layout.yaxis.dtick
                if (hi-lo)/step<=150:ax.yaxis.set_major_locator(MultipleLocator(step))
            if fig.layout.yaxis.rangemode=='tozero':ax.set_ylim(bottom=0)
            if fig.layout.xaxis.rangemode=='tozero' and not date_axis:ax.set_xlim(left=0)
            for axis,layout in [(ax.xaxis,fig.layout.xaxis),(ax.yaxis,fig.layout.yaxis)]:
                axis.grid(layout.showgrid is not False,color='#e4e9ef',linewidth=.7)
            if fig.layout.xaxis.range is not None:ax.set_xlim(*fig.layout.xaxis.range)
            if fig.layout.yaxis.range is not None:ax.set_ylim(*fig.layout.yaxis.range)
            ax.set_axisbelow(True);ax.spines[['top','right']].set_visible(False)
            if legend:
                handles,labels=ax.get_legend_handles_labels()
                if secondary is not None:
                    h,l=secondary.get_legend_handles_labels();handles+=h;labels+=l
                labels=['\n'.join(textwrap.wrap(str(v),width=46 if width_mm>=170 else 28)) for v in labels]
                if handles:f.legend(handles,labels,loc='outside lower center',ncol=ncols,fontsize=7,frameon=False)
            if date_axis:
                ax.xaxis.set_major_locator(AutoDateLocator(minticks=3,maxticks=7));ax.xaxis.set_major_formatter(DateFormatter('%m.%Y'))
            buf=io.BytesIO();f.savefig(buf,format=fmt,dpi=dpi)
            return buf.getvalue()
        finally:plt.close(f)


def export_zip(charts,tables,formats=('svg','pdf'),dpi=300,width_mm=220,metadata=None,progress=None):
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as z:
        for i,(name,fig) in enumerate(charts.items()):
            for fmt in formats: z.writestr(f'charts/{i+1:03}_{safe_name(name)}.{fmt}',figure_bytes(fig,fmt,dpi,width_mm))
            if progress: progress((i+1)/max(1,len(charts)))
        for i,(name,df) in enumerate(tables.items()): z.writestr(f'tables/{i+1:03}_{safe_name(name)}.csv',csv_bytes(df))
        if tables:
            if any(len(d)>1_048_575 for d in tables.values()):
                z.writestr('XLSX_LIMIT.txt','Таблица больше лимита строк Excel. Полные данные сохранены в CSV.')
            else: z.writestr('results.xlsx',xlsx_bytes(tables))
        z.writestr('parameters.json',json.dumps(metadata or {},ensure_ascii=False,indent=2,default=str))
    return buffer.getvalue()
